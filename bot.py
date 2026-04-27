import asyncio
import logging
import os
import re
import random
import sys
import traceback
import tempfile
import zipfile
import tarfile
import io
import importlib.util
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, List, Tuple, Optional, Any

from aiogram import Bot, Dispatcher, F, types
from aiogram.filters import Command
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.utils.keyboard import InlineKeyboardBuilder, ReplyKeyboardBuilder
from aiogram.enums import ParseMode
from aiogram.exceptions import TelegramBadRequest
from aiogram.types import FSInputFile
from telethon import TelegramClient, errors, events
from telethon.tl.functions.channels import JoinChannelRequest
from telethon.tl.functions.messages import SendReactionRequest
from telethon.tl.types import Channel, Chat, User, ReactionEmoji
from telethon.errors import FloodWaitError

# Попытка импорта библиотек для извлечения текста
try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

# Настройка логирования
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    handlers=[
        logging.StreamHandler(),
        logging.FileHandler('bot.log', encoding='utf-8')
    ]
)
logger = logging.getLogger(__name__)

# Конфигурация из переменных окружения
BOT_TOKEN = os.getenv("BOT_TOKEN")
if not BOT_TOKEN:
    raise ValueError("BOT_TOKEN не найден в переменных окружения!")

# Константы
API_ID = 32480523
API_HASH = "147839735c9fa4e83451209e9b55cfc5"
SUPPORT_USERNAME = "@VestSupport"
ADMIN_ID = 7973988177
BUY_ACCOUNTS = "@v3estnikov"
DONATION_CHANNEL = "@VestSoftTG"

# Инициализация бота и диспетчера
bot = Bot(token=BOT_TOKEN)
dp = Dispatcher(storage=MemoryStorage())

# ============ ХРАНИЛИЩА ДАННЫХ ============

# Аккаунты пользователей: user_id -> {phone: {"client": TelegramClient, "phone": str}}
user_sessions: Dict[int, Dict[str, dict]] = {}

# Активные рассылки: user_id -> {broadcast_id: Task}
active_broadcasts: Dict[int, Dict[str, asyncio.Task]] = {}

# Активные масс-реакции: user_id -> {reaction_id: {task, chats, reaction, type, delay}}
active_reactions: Dict[int, Dict[str, Dict]] = {}

# Ожидающие подтверждения логины: user_id -> {client, phone, phone_code_hash}
pending_logins: Dict[int, Dict] = {}

# Кеш диалогов: user_id -> {phone: [(chat_id, chat_name), ...]}
dialogs_cache: Dict[int, Dict[str, List[Tuple[str, str]]]] = {}

# Сообщения со списком чатов для удаления: user_id -> [message_id, ...]
user_chat_messages: Dict[int, List[int]] = {}

# Выбранный аккаунт пользователя: user_id -> phone
user_selected_account: Dict[int, str] = {}

# Команды пользователя: user_id -> {".команда": "ответ"}
user_commands: Dict[int, Dict[str, str]] = {}

# Включено ли приветствие: user_id -> bool
user_welcome_enabled: Dict[int, bool] = {}

# Активные обработчики автоответчика: user_id -> {handler_name: handler}
active_auto_reply_handlers: Dict[int, Dict] = {}

# Сохраненные конфиги: user_id -> {config_name: {commands, welcome_enabled}}
saved_configs: Dict[int, Dict[str, Dict]] = {}

# Загруженные плагины: user_id -> {plugin_name: {module, instance, active, start_time, code}}
loaded_plugins: Dict[int, Dict[str, Dict[str, Any]]] = {}

# ============ КОНСТАНТЫ ============

# Базовые реакции для масс-реакций
BASE_REACTIONS = [
    ("👍", "👍"),
    ("👎", "👎"),
    ("❤", "❤"),
    ("🔥", "🔥"),
    ("🥰", "🥰"),
    ("😁", "😁"),
]

# Юзернейм бота для проверки спам-блока
SPAM_BOT_USERNAME = "@spambot"

# Плейсхолдеры для подстановки в командах
PLACEHOLDERS = {
    "{NICK}": "first_name",
    "{FIRSTNAME}": "first_name",
    "{LASTNAME}": "last_name",
    "{FULLNAME}": "full_name",
    "{USERNAME}": "username",
    "{ID}": "id",
    "{PHONE}": "phone",
}

# ============ СОСТОЯНИЯ FSM ============

class AccountStates(StatesGroup):
    """Состояния для добавления аккаунта"""
    waiting_for_phone = State()
    waiting_for_code = State()
    waiting_for_2fa = State()

class BroadcastStates(StatesGroup):
    """Состояния для создания рассылки"""
    adding_messages = State()
    waiting_for_count = State()
    waiting_for_delay = State()
    selecting_chats = State()
    selecting_mode = State()

class JoinStates(StatesGroup):
    """Состояния для вступления в чаты"""
    waiting_for_usernames = State()

class AdminStates(StatesGroup):
    """Состояния для админ-панели"""
    waiting_for_broadcast = State()
    waiting_for_confirm = State()

class ReactionStates(StatesGroup):
    """Состояния для масс-реакций"""
    selecting_chats = State()
    waiting_for_delay = State()
    selecting_type = State()
    selecting_reaction = State()

class CommandStates(StatesGroup):
    """Состояния для добавления команд"""
    waiting_for_command_name = State()
    waiting_for_command_response = State()

class WelcomeStates(StatesGroup):
    """Состояния для настройки приветствия"""
    waiting_for_welcome_message = State()

class ConfigStates(StatesGroup):
    """Состояния для сохранения конфига"""
    waiting_for_config_name = State()

class PluginStates(StatesGroup):
    """Состояния для загрузки плагина"""
    waiting_for_plugin_file = State()
    waiting_for_plugin_name = State()

# ============ ПРЕМИУМ ЭМОДЗИ ДЛЯ СООБЩЕНИЙ ============

E = {
    "settings": '<tg-emoji emoji-id="5870982283724328568">⚙</tg-emoji>',
    "profile": '<tg-emoji emoji-id="5870994129244131212">👤</tg-emoji>',
    "people": '<tg-emoji emoji-id="5870772616305839506">👥</tg-emoji>',
    "file": '<tg-emoji emoji-id="5870528606328852614">📁</tg-emoji>',
    "stats": '<tg-emoji emoji-id="5870930636742595124">📊</tg-emoji>',
    "home": '<tg-emoji emoji-id="5873147866364514353">🏘</tg-emoji>',
    "lock": '<tg-emoji emoji-id="6037249452824072506">🔒</tg-emoji>',
    "unlock": '<tg-emoji emoji-id="6037496202990194718">🔓</tg-emoji>',
    "megaphone": '<tg-emoji emoji-id="6039422865189638057">📣</tg-emoji>',
    "check": '<tg-emoji emoji-id="5870633910337015697">✅</tg-emoji>',
    "cross": '<tg-emoji emoji-id="5870657884844462243">❌</tg-emoji>',
    "trash": '<tg-emoji emoji-id="5870875489362513438">🗑</tg-emoji>',
    "back": '<tg-emoji emoji-id="5893057118545646106">◁</tg-emoji>',
    "link": '<tg-emoji emoji-id="5769289093221454192">🔗</tg-emoji>',
    "info": '<tg-emoji emoji-id="6028435952299413210">ℹ</tg-emoji>',
    "bot": '<tg-emoji emoji-id="6030400221232501136">🤖</tg-emoji>',
    "eye": '<tg-emoji emoji-id="6037397706505195857">👁</tg-emoji>',
    "send": '<tg-emoji emoji-id="5963103826075456248">⬆</tg-emoji>',
    "gift": '<tg-emoji emoji-id="6032644646587338669">🎁</tg-emoji>',
    "clock": '<tg-emoji emoji-id="5983150113483134607">⏰</tg-emoji>',
    "celebration": '<tg-emoji emoji-id="6041731551845159060">🎉</tg-emoji>',
    "write": '<tg-emoji emoji-id="5870753782874246579">✍</tg-emoji>',
    "apps": '<tg-emoji emoji-id="5778672437122045013">📦</tg-emoji>',
    "code": '<tg-emoji emoji-id="5940433880585605708">🔨</tg-emoji>',
    "loading": '<tg-emoji emoji-id="5345906554510012647">🔄</tg-emoji>',
    "user_check": '<tg-emoji emoji-id="5891207662678317861">👤</tg-emoji>',
    "wallet": '<tg-emoji emoji-id="5769126056262898415">👛</tg-emoji>',
    "notification": '<tg-emoji emoji-id="6039486778597970865">🔔</tg-emoji>',
    "smile": '<tg-emoji emoji-id="5870764288364252592">🙂</tg-emoji>',
    "download": '<tg-emoji emoji-id="6039802767931871481">⬇</tg-emoji>',
    "money": '<tg-emoji emoji-id="5904462880941545555">🪙</tg-emoji>',
    "box": '<tg-emoji emoji-id="5884479287171485878">📦</tg-emoji>',
    "pencil": '<tg-emoji emoji-id="5870676941614354370">🖋</tg-emoji>',
    "paperclip": '<tg-emoji emoji-id="6039451237743595514">📎</tg-emoji>',
    "eye_hidden": '<tg-emoji emoji-id="6037243349675544634">👁</tg-emoji>',
    "font": '<tg-emoji emoji-id="5870801517140775623">🔗</tg-emoji>',
    "media": '<tg-emoji emoji-id="6035128606563241721">🖼</tg-emoji>',
    "geo": '<tg-emoji emoji-id="6042011682497106307">📍</tg-emoji>',
    "calendar": '<tg-emoji emoji-id="5890937706803894250">📅</tg-emoji>',
    "tag": '<tg-emoji emoji-id="5886285355279193209">🏷</tg-emoji>',
    "time_passed": '<tg-emoji emoji-id="5775896410780079073">🕓</tg-emoji>',
    "brush": '<tg-emoji emoji-id="6050679691004612757">🖌</tg-emoji>',
    "add_text": '<tg-emoji emoji-id="5771851822897566479">🔡</tg-emoji>',
    "resolution": '<tg-emoji emoji-id="5778479949572738874">↔</tg-emoji>',
    "money_send": '<tg-emoji emoji-id="5890848474563352982">🪙</tg-emoji>',
    "money_accept": '<tg-emoji emoji-id="5879814368572478751">🏧</tg-emoji>',
}

# ============ ID ЭМОДЗИ ДЛЯ КНОПОК ============

E_ID = {
    "settings": "5870982283724328568",
    "profile": "5870994129244131212",
    "people": "5870772616305839506",
    "file": "5870528606328852614",
    "stats": "5870930636742595124",
    "home": "5873147866364514353",
    "lock": "6037249452824072506",
    "unlock": "6037496202990194718",
    "megaphone": "6039422865189638057",
    "check": "5870633910337015697",
    "cross": "5870657884844462243",
    "trash": "5870875489362513438",
    "back": "5893057118545646106",
    "link": "5769289093221454192",
    "info": "6028435952299413210",
    "bot": "6030400221232501136",
    "eye": "6037397706505195857",
    "send": "5963103826075456248",
    "gift": "6032644646587338669",
    "clock": "5983150113483134607",
    "celebration": "6041731551845159060",
    "write": "5870753782874246579",
    "apps": "5778672437122045013",
    "code": "5940433880585605708",
    "loading": "5345906554510012647",
    "user_check": "5891207662678317861",
    "wallet": "5769126056262898415",
    "notification": "6039486778597970865",
    "smile": "5870764288364252592",
    "download": "6039802767931871481",
    "money": "5904462880941545555",
    "box": "5884479287171485878",
    "pencil": "5870676941614354370",
    "subscribe": "6039450962865688331",
    "check_sub": "5774022692642492953",
    "blue": "5373141891321699086",
    "red": "5370810157871667232",
    "green": "5471984997361523302",
}

# ============ ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ ============

def em(key: str) -> str:
    """Возвращает полный тег премиум эмодзи с символом для HTML"""
    return E.get(key, "")

def eid(key: str) -> str:
    """Возвращает ID эмодзи для использования в кнопках"""
    return E_ID.get(key, "")

def process_placeholders(text: str, user_info: dict) -> str:
    """Заменяет плейсхолдеры {NICK}, {FIRSTNAME} и т.д. на реальные данные пользователя"""
    result = text
    for placeholder, field in PLACEHOLDERS.items():
        if placeholder in result:
            value = user_info.get(field, "")
            if field == "full_name":
                first = user_info.get("first_name", "")
                last = user_info.get("last_name", "")
                value = f"{first} {last}".strip() or "Пользователь"
            result = result.replace(placeholder, str(value) if value else placeholder)
    return result

# ============ КЛАВИАТУРЫ ============

def get_main_keyboard():
    """Главное меню с ReplyKeyboard"""
    builder = ReplyKeyboardBuilder()
    builder.button(
        text="Менеджер аккаунтов",
        icon_custom_emoji_id=eid("settings")
    )
    builder.button(
        text="Функции",
        icon_custom_emoji_id=eid("stats")
    )
    builder.button(
        text="Поддержка",
        icon_custom_emoji_id=eid("megaphone")
    )
    builder.adjust(2, 1)
    return builder.as_markup(resize_keyboard=True)

def get_accounts_menu_keyboard():
    """Меню управления аккаунтами"""
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Добавить аккаунт",
        callback_data="add_account",
        style="primary",
        icon_custom_emoji_id=eid("gift")
    )
    builder.button(
        text="Мои аккаунты",
        callback_data="my_accounts",
        style="success",
        icon_custom_emoji_id=eid("profile")
    )
    builder.button(
        text="Выбрать аккаунт",
        callback_data="select_account",
        style="default",
        icon_custom_emoji_id=eid("user_check")
    )
    builder.button(
        text="Назад",
        callback_data="main_menu",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(2, 1)
    return builder.as_markup()

def get_functions_keyboard():
    """Меню функций бота"""
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Рассылка",
        callback_data="broadcast",
        style="primary",
        icon_custom_emoji_id=eid("megaphone")
    )
    builder.button(
        text="Вступление в чаты",
        callback_data="join_chats",
        style="success",
        icon_custom_emoji_id=eid("people")
    )
    builder.button(
        text="Масс реакции",
        callback_data="mass_reactions",
        style="primary",
        icon_custom_emoji_id=eid("smile")
    )
    builder.button(
        text="Проверка спам блока",
        callback_data="check_spam_block",
        style="danger",
        icon_custom_emoji_id=eid("eye")
    )
    builder.button(
        text="Команды",
        callback_data="commands_menu",
        style="success",
        icon_custom_emoji_id=eid("code")
    )
    builder.button(
        text="Плагины",
        callback_data="plugins_menu",
        style="primary",
        icon_custom_emoji_id=eid("apps")
    )
    builder.button(
        text="Назад",
        callback_data="main_menu",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(1)
    return builder.as_markup()

def get_plugins_menu_keyboard():
    """Меню плагинов"""
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Загрузить плагин",
        callback_data="upload_plugin",
        style="primary",
        icon_custom_emoji_id=eid("send")
    )
    builder.button(
        text="Мои плагины",
        callback_data="my_plugins",
        style="success",
        icon_custom_emoji_id=eid("box")
    )
    builder.button(
        text="Документация",
        callback_data="plugin_docs",
        style="default",
        icon_custom_emoji_id=eid("file")
    )
    builder.button(
        text="Назад",
        callback_data="main_menu",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(1)
    return builder.as_markup()

def get_plugin_actions_keyboard(plugin_name: str, is_active: bool):
    """Клавиатура управления конкретным плагином"""
    builder = InlineKeyboardBuilder()
    if is_active:
        builder.button(
            text="Остановить",
            callback_data=f"plg_stop_{plugin_name}",
            style="danger",
            icon_custom_emoji_id=eid("cross")
        )
    else:
        builder.button(
            text="Запустить",
            callback_data=f"plg_start_{plugin_name}",
            style="success",
            icon_custom_emoji_id=eid("send")
        )
    builder.button(
        text="Статистика",
        callback_data=f"plg_stats_{plugin_name}",
        style="default",
        icon_custom_emoji_id=eid("stats")
    )
    builder.button(
        text="Удалить",
        callback_data=f"plg_del_{plugin_name}",
        style="danger",
        icon_custom_emoji_id=eid("trash")
    )
    builder.button(
        text="Назад",
        callback_data="my_plugins",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(2, 2, 1)
    return builder.as_markup()

def get_commands_keyboard():
    """Меню команд"""
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Добавить команду",
        callback_data="add_command",
        style="primary",
        icon_custom_emoji_id=eid("code")
    )
    builder.button(
        text="Мои команды",
        callback_data="my_commands",
        style="success",
        icon_custom_emoji_id=eid("file")
    )
    builder.button(
        text="Приветствие новым",
        callback_data="welcome_setup",
        style="primary",
        icon_custom_emoji_id=eid("gift")
    )
    builder.button(
        text="Сохранить конфиг",
        callback_data="save_config",
        style="default",
        icon_custom_emoji_id=eid("download")
    )
    builder.button(
        text="Загрузить конфиг",
        callback_data="load_config",
        style="default",
        icon_custom_emoji_id=eid("box")
    )
    builder.button(
        text="Назад",
        callback_data="main_menu",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(1)
    return builder.as_markup()

def get_admin_keyboard():
    """Меню админ-панели"""
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Статистика",
        callback_data="admin_stats",
        style="primary",
        icon_custom_emoji_id=eid("stats")
    )
    builder.button(
        text="Рассылка всем",
        callback_data="admin_broadcast",
        style="success",
        icon_custom_emoji_id=eid("megaphone")
    )
    builder.button(
        text="Назад",
        callback_data="main_menu",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(1)
    return builder.as_markup()

def get_back_keyboard(callback_data: str = "main_menu"):
    """Клавиатура с кнопкой Назад"""
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Назад",
        callback_data=callback_data,
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    return builder.as_markup()

def get_mode_keyboard():
    """Клавиатура выбора режима рассылки"""
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Одновременная",
        callback_data="mode_sync",
        style="primary",
        icon_custom_emoji_id=eid("send")
    )
    builder.button(
        text="Рандомная",
        callback_data="mode_random",
        style="success",
        icon_custom_emoji_id=eid("loading")
    )
    builder.button(
        text="Назад",
        callback_data="cancel_broadcast",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(2, 1)
    return builder.as_markup()

def get_message_actions_keyboard():
    """Клавиатура действий после добавления сообщения"""
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Добавить сообщение",
        callback_data="add_msg",
        style="primary",
        icon_custom_emoji_id=eid("write")
    )
    builder.button(
        text="Запустить рассылку",
        callback_data="start_msg_config",
        style="success",
        icon_custom_emoji_id=eid("send")
    )
    builder.button(
        text="Отмена",
        callback_data="cancel_broadcast",
        style="danger",
        icon_custom_emoji_id=eid("cross")
    )
    builder.adjust(1)
    return builder.as_markup()

def get_reaction_type_keyboard():
    """Клавиатура выбора типа масс-реакций"""
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Всё подряд",
        callback_data="rtype_all",
        style="primary",
        icon_custom_emoji_id=eid("send")
    )
    builder.button(
        text="Через один",
        callback_data="rtype_skip",
        style="success",
        icon_custom_emoji_id=eid("loading")
    )
    builder.button(
        text="Отмена",
        callback_data="cancel_reactions",
        style="danger",
        icon_custom_emoji_id=eid("cross")
    )
    builder.adjust(2, 1)
    return builder.as_markup()

def get_reactions_keyboard():
    """Клавиатура выбора реакции"""
    builder = InlineKeyboardBuilder()
    for emoji_text, _ in BASE_REACTIONS:
        builder.button(
            text=emoji_text,
            callback_data="react_" + emoji_text,
            style="default"
        )
    builder.button(
        text="Отмена",
        callback_data="cancel_reactions",
        style="danger",
        icon_custom_emoji_id=eid("cross")
    )
    builder.adjust(3, 3, 1)
    return builder.as_markup()

# ============ УТИЛИТЫ ДЛЯ ОТПРАВКИ СООБЩЕНИЙ ============

async def safe_send(chat_id: int, text: str, reply_markup=None, **kwargs):
    """Безопасная отправка сообщения с fallback при ошибке парсинга HTML"""
    try:
        return await bot.send_message(
            chat_id,
            text,
            parse_mode=ParseMode.HTML,
            reply_markup=reply_markup,
            **kwargs
        )
    except TelegramBadRequest as e:
        if "can't parse" in str(e):
            # Удаляем все HTML теги и премиум эмодзи
            clean = re.sub(r'<tg-emoji[^>]*>.*?</tg-emoji>', '', text)
            clean = re.sub(r'<[^>]+>', '', clean)
            return await bot.send_message(
                chat_id,
                clean,
                reply_markup=reply_markup,
                **kwargs
            )
        raise

async def safe_edit(message: types.Message, text: str, reply_markup=None):
    """Безопасное редактирование сообщения с fallback на новое сообщение"""
    try:
        return await message.edit_text(
            text,
            parse_mode=ParseMode.HTML,
            reply_markup=reply_markup
        )
    except TelegramBadRequest as e:
        err = str(e)
        if "message can't be edited" in err or "message is not modified" in err:
            try:
                await message.delete()
            except:
                pass
            return await safe_send(message.chat.id, text, reply_markup=reply_markup)
        if "can't parse" in err:
            try:
                await message.delete()
            except:
                pass
            clean = re.sub(r'<tg-emoji[^>]*>.*?</tg-emoji>', '', text)
            clean = re.sub(r'<[^>]+>', '', clean)
            return await safe_send(message.chat.id, clean, reply_markup=reply_markup)
        raise
    except Exception:
        try:
            await message.delete()
        except:
            pass
        return await safe_send(message.chat.id, text, reply_markup=reply_markup)

async def delete_user_chat_messages(user_id: int):
    """Удаляет все сохраненные сообщения со списком чатов"""
    if user_id in user_chat_messages:
        for msg_id in user_chat_messages[user_id]:
            try:
                await bot.delete_message(user_id, msg_id)
            except:
                pass
        user_chat_messages[user_id] = []

async def add_user_chat_message(user_id: int, message_id: int):
    """Добавляет ID сообщения в список для последующего удаления"""
    if user_id not in user_chat_messages:
        user_chat_messages[user_id] = []
    user_chat_messages[user_id].append(message_id)

def get_active_account(user_id: int) -> Optional[Tuple[str, TelegramClient]]:
    """Возвращает выбранный аккаунт или первый доступный"""
    accounts = user_sessions.get(user_id, {})
    if not accounts:
        return None
    
    selected_phone = user_selected_account.get(user_id)
    if selected_phone and selected_phone in accounts:
        return selected_phone, accounts[selected_phone]["client"]
    
    # Если нет выбранного - берем первый
    phone = list(accounts.keys())[0]
    return phone, accounts[phone]["client"]

# ============ ИЗВЛЕЧЕНИЕ КОДА ИЗ ФАЙЛОВ ============

async def extract_code_from_file(file_data: bytes, filename: str) -> Optional[str]:
    """Извлекает Python код из файла любого поддерживаемого формата"""
    ext = filename.lower().split('.')[-1] if '.' in filename else ''
    
    # Обработка Python файлов
    if ext in ['py', 'pyw']:
        return file_data.decode('utf-8', errors='ignore')
    
    # Обработка текстовых файлов
    if ext in ['txt', 'cfg', 'ini', 'env', 'md']:
        try:
            text = file_data.decode('utf-8')
            # Ищем код в Python блоках
            python_blocks = re.findall(r'```python\n(.*?)```', text, re.DOTALL)
            if python_blocks:
                return '\n\n'.join(python_blocks)
            # Проверяем наличие Python синтаксиса
            if re.search(r'(import |def |class |async def |from .+ import)', text):
                return text
        except:
            pass
    
    # Обработка PDF
    if ext == 'pdf':
        pdf_text = await extract_text_from_pdf(file_data)
        if pdf_text:
            python_blocks = re.findall(r'```python\n(.*?)```', pdf_text, re.DOTALL)
            if python_blocks:
                return '\n\n'.join(python_blocks)
            if re.search(r'(import |def |class |async def |from .+ import)', pdf_text):
                return pdf_text
        return None
    
    # Обработка DOCX
    if ext == 'docx':
        docx_text = await extract_text_from_docx(file_data)
        if docx_text:
            python_blocks = re.findall(r'```python\n(.*?)```', docx_text, re.DOTALL)
            if python_blocks:
                return '\n\n'.join(python_blocks)
            if re.search(r'(import |def |class |async def |from .+ import)', docx_text):
                return docx_text
        return None
    
    # Обработка архивов
    if ext in ['zip', 'tar', 'gz', 'bz2', 'tar.gz', 'tar.bz2']:
        try:
            # Сохраняем архив во временный файл
            with tempfile.NamedTemporaryFile(delete=False, suffix='.' + ext) as tmp:
                tmp.write(file_data)
                tmp_path = tmp.name
            
            extract_dir = tempfile.mkdtemp()
            
            # Распаковываем
            if ext == 'zip':
                with zipfile.ZipFile(tmp_path, 'r') as zf:
                    zf.extractall(extract_dir)
            elif ext in ['tar', 'gz', 'bz2', 'tar.gz', 'tar.bz2']:
                with tarfile.open(tmp_path, 'r:*') as tf:
                    tf.extractall(extract_dir)
            
            # Ищем Python файлы
            code_parts = []
            for root, dirs, files in os.walk(extract_dir):
                for f in sorted(files):
                    ext_f = f.split('.')[-1] if '.' in f else ''
                    if ext_f in ['py', 'pyw']:
                        fpath = os.path.join(root, f)
                        with open(fpath, 'r', encoding='utf-8', errors='ignore') as pf:
                            code_parts.append(f"# --- {f} ---\n{pf.read()}")
                    elif ext_f in ['txt', 'cfg', 'ini']:
                        fpath = os.path.join(root, f)
                        with open(fpath, 'r', encoding='utf-8', errors='ignore') as pf:
                            text = pf.read()
                            python_blocks = re.findall(r'```python\n(.*?)```', text, re.DOTALL)
                            if python_blocks:
                                code_parts.extend(python_blocks)
                            elif re.search(r'(import |def |class |async def )', text):
                                code_parts.append(f"# --- {f} ---\n{text}")
            
            # Очищаем временные файлы
            os.unlink(tmp_path)
            import shutil
            shutil.rmtree(extract_dir, ignore_errors=True)
            
            if code_parts:
                return '\n\n'.join(code_parts)
        except Exception as ex:
            logger.error(f"Error extracting archive: {ex}")
    
    # Пробуем декодировать как обычный текст
    try:
        text = file_data.decode('utf-8')
        python_blocks = re.findall(r'```python\n(.*?)```', text, re.DOTALL)
        if python_blocks:
            return '\n\n'.join(python_blocks)
        if re.search(r'(import |def |class |async def |from .+ import)', text):
            return text
    except:
        pass
    
    return None

async def extract_text_from_pdf(file_data: bytes) -> Optional[str]:
    """Извлекает текст из PDF файла используя PyPDF2 или pdfplumber"""
    text_parts = []
    
    # Пробуем PyPDF2
    if PyPDF2:
        try:
            pdf_file = io.BytesIO(file_data)
            reader = PyPDF2.PdfReader(pdf_file)
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
            if text_parts:
                return '\n\n'.join(text_parts)
        except Exception as ex:
            logger.debug(f"PyPDF2 extraction failed: {ex}")
    
    # Пробуем pdfplumber
    if pdfplumber:
        try:
            pdf_file = io.BytesIO(file_data)
            with pdfplumber.open(pdf_file) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)
            if text_parts:
                return '\n\n'.join(text_parts)
        except Exception as ex:
            logger.debug(f"pdfplumber extraction failed: {ex}")
    
    return None

async def extract_text_from_docx(file_data: bytes) -> Optional[str]:
    """Извлекает текст из DOCX файла"""
    if DocxDocument:
        try:
            docx_file = io.BytesIO(file_data)
            doc = DocxDocument(docx_file)
            text_parts = [para.text for para in doc.paragraphs if para.text]
            if text_parts:
                return '\n'.join(text_parts)
        except Exception as ex:
            logger.debug(f"DOCX extraction failed: {ex}")
    return None

def validate_plugin_code(code: str) -> Tuple[bool, str]:
    """Проверяет, содержит ли код необходимые элементы плагина"""
    required_elements = [
        ('class Plugin', 'Класс Plugin'),
        ('async def start', 'Метод start'),
        ('async def stop', 'Метод stop'),
    ]
    
    for element, name in required_elements:
        if element not in code:
            return False, f"Отсутствует: {name}"
    
    return True, "OK"

async def load_plugin_module(code: str, plugin_name: str) -> Any:
    """Загружает плагин как модуль Python из строки кода"""
    plugin_dir = Path("plugins")
    plugin_dir.mkdir(exist_ok=True)
    
    plugin_file = plugin_dir / f"{plugin_name}.py"
    plugin_file.write_text(code, encoding='utf-8')
    
    spec = importlib.util.spec_from_file_location(plugin_name, plugin_file)
    module = importlib.util.module_from_spec(spec)
    sys.modules[plugin_name] = module
    spec.loader.exec_module(module)
    
    return module

# ============ КОМАНДА START ============

@dp.message(Command("start"))
async def cmd_start(message: types.Message):
    """Обработчик команды /start"""
    text = (
        em('bot') + " <b>Добро пожаловать!</b>\n\n"
        "<b>Главное меню:</b>\n"
        + em('settings') + " <b>Менеджер аккаунтов</b> — управление аккаунтами\n"
        + em('stats') + " <b>Функции</b> — рассылка, команды, реакции, плагины\n"
        + em('megaphone') + " <b>Поддержка</b> — связь с поддержкой\n\n"
        + em('wallet') + " <b>Купить аккаунт для рассылки:</b> " + BUY_ACCOUNTS + "\n"
        + em('link') + " <b>Документация по плагинам:</b> " + DONATION_CHANNEL
    )
    await safe_send(message.chat.id, text, reply_markup=get_main_keyboard())

@dp.message(Command("admin"))
async def admin_command(message: types.Message):
    """Обработчик команды /admin"""
    if message.from_user.id != ADMIN_ID:
        await safe_send(message.chat.id, em('cross') + " Нет доступа")
        return
    
    text = (
        em('lock') + " <b>Админ панель</b>\n\n"
        + em('stats') + " <b>Статистика</b> — просмотр статистики\n"
        + em('megaphone') + " <b>Рассылка всем</b> — отправить сообщение всем пользователям"
    )
    await safe_send(message.chat.id, text, reply_markup=get_admin_keyboard())

# ============ ОБРАБОТЧИКИ ТЕКСТОВЫХ КНОПОК ============

@dp.message(F.text == "Менеджер аккаунтов")
async def accounts_manager(message: types.Message):
    """Обработчик кнопки Менеджер аккаунтов"""
    user_id = message.from_user.id
    count = len(user_sessions.get(user_id, {}))
    selected = user_selected_account.get(user_id, "Не выбран")
    
    text = (
        em('settings') + " <b>Менеджер аккаунтов</b>\n"
        + em('profile') + " Активных аккаунтов: " + str(count) + " (безлимит)\n"
        + em('user_check') + " Выбран: <code>" + selected + "</code>"
    )
    await safe_send(message.chat.id, text, reply_markup=get_accounts_menu_keyboard())

@dp.message(F.text == "Функции")
async def functions_menu(message: types.Message):
    """Обработчик кнопки Функции"""
    text = (
        em('stats') + " <b>Функции</b>\n"
        + em('megaphone') + " Рассылка сообщений\n"
        + em('people') + " Вступление в чаты\n"
        + em('smile') + " Масс реакции\n"
        + em('eye') + " Проверка спам блока\n"
        + em('code') + " Команды\n"
        + em('apps') + " Плагины"
    )
    await safe_send(message.chat.id, text, reply_markup=get_functions_keyboard())

@dp.message(F.text == "Поддержка")
async def support(message: types.Message):
    """Обработчик кнопки Поддержка"""
    text = (
        em('megaphone') + " <b>Поддержка</b>\n"
        + em('link') + " Свяжитесь с нами: " + SUPPORT_USERNAME
    )
    await safe_send(message.chat.id, text, reply_markup=get_main_keyboard())

@dp.callback_query(F.data == "main_menu")
async def back_to_main(callback: types.CallbackQuery):
    """Возврат в главное меню"""
    try:
        await callback.message.delete()
    except:
        pass
    
    text = em('bot') + " <b>Главное меню</b>"
    await safe_send(callback.message.chat.id, text, reply_markup=get_main_keyboard())
    await callback.answer()

# ============ ДОБАВЛЕНИЕ АККАУНТА ============

@dp.callback_query(F.data == "add_account")
async def add_account(callback: types.CallbackQuery, state: FSMContext):
    """Начало процесса добавления аккаунта"""
    text = (
        em('apps') + " <b>Добавление аккаунта</b>\n"
        + em('write') + " Введите номер телефона: <code>+79123456789</code>"
    )
    await safe_edit(callback.message, text, reply_markup=get_back_keyboard("accounts_manager"))
    await state.set_state(AccountStates.waiting_for_phone)
    await callback.answer()

@dp.message(AccountStates.waiting_for_phone)
async def process_phone(message: types.Message, state: FSMContext):
    """Обработка введенного номера телефона"""
    phone = message.text.strip()
    
    # Валидация номера
    if not re.match(r'^\+\d{10,15}$', phone):
        await safe_send(
            message.chat.id,
            em('cross') + " Неверный формат. Пример: <code>+79123456789</code>",
            reply_markup=get_back_keyboard("accounts_manager")
        )
        return
    
    user_id = message.from_user.id
    
    # Проверка на дубликат
    if user_id in user_sessions and phone in user_sessions[user_id]:
        await safe_send(
            message.chat.id,
            em('cross') + " Этот аккаунт уже добавлен!",
            reply_markup=get_accounts_menu_keyboard()
        )
        await state.clear()
        return
    
    # Создание клиента Telethon
    client = TelegramClient(
        'sessions/' + str(user_id) + '_' + phone.replace("+", ""),
        API_ID,
        API_HASH
    )
    
    try:
        await client.connect()
        sent_code = await client.send_code_request(phone)
        
        # Сохраняем данные для подтверждения
        pending_logins[user_id] = {
            "client": client,
            "phone": phone,
            "phone_code_hash": sent_code.phone_code_hash
        }
        
        text = (
            em('gift') + " Код отправлен на <code>" + phone + "</code>\n"
            + em('write') + " Введите код из SMS:"
        )
        await safe_send(
            message.chat.id,
            text,
            reply_markup=get_back_keyboard("accounts_manager")
        )
        await state.set_state(AccountStates.waiting_for_code)
        
    except Exception as ex:
        await client.disconnect()
        await safe_send(
            message.chat.id,
            em('cross') + " Ошибка: " + str(ex),
            reply_markup=get_accounts_menu_keyboard()
        )
        await state.clear()

@dp.message(AccountStates.waiting_for_code)
async def process_code(message: types.Message, state: FSMContext):
    """Обработка кода подтверждения"""
    code = message.text.strip()
    user_id = message.from_user.id
    
    if user_id not in pending_logins:
        await safe_send(
            message.chat.id,
            em('cross') + " Сессия истекла",
            reply_markup=get_accounts_menu_keyboard()
        )
        await state.clear()
        return
    
    data = pending_logins[user_id]
    client = data["client"]
    phone = data["phone"]
    phone_code_hash = data["phone_code_hash"]
    
    try:
        await client.sign_in(phone=phone, code=code, phone_code_hash=phone_code_hash)
        await on_successful_login(user_id, phone, client, message, state)
    except errors.SessionPasswordNeededError:
        # Требуется 2FA
        text = (
            em('lock') + " Требуется двухфакторная аутентификация\n"
            + em('write') + " Введите пароль 2FA:"
        )
        await safe_send(
            message.chat.id,
            text,
            reply_markup=get_back_keyboard("accounts_manager")
        )
        await state.set_state(AccountStates.waiting_for_2fa)
    except Exception as ex:
        await client.disconnect()
        pending_logins.pop(user_id, None)
        await safe_send(
            message.chat.id,
            em('cross') + " Ошибка: " + str(ex),
            reply_markup=get_accounts_menu_keyboard()
        )
        await state.clear()

@dp.message(AccountStates.waiting_for_2fa)
async def process_2fa(message: types.Message, state: FSMContext):
    """Обработка пароля 2FA"""
    password = message.text.strip()
    user_id = message.from_user.id
    
    if user_id not in pending_logins:
        await safe_send(
            message.chat.id,
            em('cross') + " Сессия истекла",
            reply_markup=get_accounts_menu_keyboard()
        )
        await state.clear()
        return
    
    data = pending_logins[user_id]
    client = data["client"]
    phone = data["phone"]
    
    try:
        await client.sign_in(password=password)
        await on_successful_login(user_id, phone, client, message, state)
    except Exception as ex:
        await client.disconnect()
        pending_logins.pop(user_id, None)
        await safe_send(
            message.chat.id,
            em('cross') + " Ошибка 2FA: " + str(ex),
            reply_markup=get_accounts_menu_keyboard()
        )
        await state.clear()

async def on_successful_login(user_id: int, phone: str, client: TelegramClient, 
                              message: types.Message, state: FSMContext):
    """Действия при успешном входе в аккаунт"""
    if user_id not in user_sessions:
        user_sessions[user_id] = {}
    
    user_sessions[user_id][phone] = {"client": client, "phone": phone}
    pending_logins.pop(user_id, None)
    dialogs_cache.pop(user_id, None)
    
    # Автоматически выбираем первый аккаунт
    if user_id not in user_selected_account:
        user_selected_account[user_id] = phone
    
    await safe_send(
        message.chat.id,
        em('check') + " Аккаунт <code>" + phone + "</code> успешно добавлен!",
        reply_markup=get_accounts_menu_keyboard()
    )
    await state.clear()

# ============ УПРАВЛЕНИЕ АККАУНТАМИ ============

@dp.callback_query(F.data == "my_accounts")
async def my_accounts(callback: types.CallbackQuery):
    """Просмотр списка аккаунтов"""
    user_id = callback.from_user.id
    accounts = user_sessions.get(user_id, {})
    
    if not accounts:
        text = (
            em('unlock') + " <b>Мои аккаунты</b>\n\n"
            + em('cross') + " Нет добавленных аккаунтов"
        )
        await safe_edit(callback.message, text, reply_markup=get_accounts_menu_keyboard())
        await callback.answer()
        return
    
    builder = InlineKeyboardBuilder()
    for phone in accounts:
        builder.button(text=phone, callback_data="acc_" + phone, style="default")
    builder.button(
        text="Назад",
        callback_data="accounts_manager",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(1)
    
    text = (
        em('profile') + " <b>Мои аккаунты</b>\n"
        + em('unlock') + " Всего: " + str(len(accounts)) + "\n"
        "Выберите аккаунт для управления:"
    )
    await safe_edit(callback.message, text, reply_markup=builder.as_markup())
    await callback.answer()

@dp.callback_query(F.data == "select_account")
async def select_account_menu(callback: types.CallbackQuery):
    """Меню выбора активного аккаунта"""
    user_id = callback.from_user.id
    accounts = user_sessions.get(user_id, {})
    
    if not accounts:
        await callback.answer("Нет добавленных аккаунтов!", show_alert=True)
        return
    
    builder = InlineKeyboardBuilder()
    for phone in accounts:
        is_selected = user_selected_account.get(user_id) == phone
        prefix = "✅ " if is_selected else ""
        builder.button(
            text=prefix + phone,
            callback_data="sel_" + phone,
            style="success" if is_selected else "default"
        )
    builder.button(
        text="Назад",
        callback_data="accounts_manager",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(1)
    
    selected = user_selected_account.get(user_id, "Не выбран")
    text = (
        em('user_check') + " <b>Выбор аккаунта</b>\n"
        "Текущий: <code>" + selected + "</code>\n\n"
        "Выберите аккаунт для работы:"
    )
    await safe_edit(callback.message, text, reply_markup=builder.as_markup())
    await callback.answer()

@dp.callback_query(F.data.startswith("sel_"))
async def select_account(callback: types.CallbackQuery):
    """Выбор аккаунта"""
    phone = callback.data.replace("sel_", "")
    user_id = callback.from_user.id
    
    if user_id in user_sessions and phone in user_sessions[user_id]:
        user_selected_account[user_id] = phone
        dialogs_cache.pop(user_id, None)  # Сбрасываем кеш диалогов
        await callback.answer("Выбран аккаунт " + phone, show_alert=True)
        await select_account_menu(callback)
    else:
        await callback.answer("Аккаунт не найден", show_alert=True)

@dp.callback_query(F.data.startswith("acc_"))
async def account_info(callback: types.CallbackQuery):
    """Информация об аккаунте"""
    phone = callback.data.replace("acc_", "")
    user_id = callback.from_user.id
    is_selected = user_selected_account.get(user_id) == phone
    
    builder = InlineKeyboardBuilder()
    if not is_selected:
        builder.button(
            text="Выбрать",
            callback_data="sel_" + phone,
            style="success",
            icon_custom_emoji_id=eid("user_check")
        )
    builder.button(
        text="Удалить",
        callback_data="del_" + phone,
        style="danger",
        icon_custom_emoji_id=eid("trash")
    )
    builder.button(
        text="Назад",
        callback_data="my_accounts",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(1)
    
    status = "✅ Выбран" if is_selected else "⚪ Доступен"
    text = (
        em('profile') + " <b>Аккаунт:</b> <code>" + phone + "</code>\n"
        + em('check') + " Статус: " + status
    )
    await safe_edit(callback.message, text, reply_markup=builder.as_markup())
    await callback.answer()

@dp.callback_query(F.data.startswith("del_"))
async def delete_account(callback: types.CallbackQuery):
    """Удаление аккаунта"""
    phone = callback.data.replace("del_", "")
    user_id = callback.from_user.id
    
    if user_id in user_sessions and phone in user_sessions[user_id]:
        try:
            await user_sessions[user_id][phone]["client"].disconnect()
        except:
            pass
        
        del user_sessions[user_id][phone]
        dialogs_cache.pop(user_id, None)
        
        # Если удалили выбранный аккаунт - выбираем следующий
        if user_selected_account.get(user_id) == phone:
            user_selected_account.pop(user_id, None)
            if user_sessions.get(user_id):
                user_selected_account[user_id] = list(user_sessions[user_id].keys())[0]
    
    await safe_edit(
        callback.message,
        em('check') + " Аккаунт <code>" + phone + "</code> удален",
        reply_markup=get_accounts_menu_keyboard()
    )
    await callback.answer("Аккаунт удален!", show_alert=True)

@dp.callback_query(F.data == "accounts_manager")
async def back_to_accounts(callback: types.CallbackQuery):
    """Возврат в меню аккаунтов"""
    user_id = callback.from_user.id
    count = len(user_sessions.get(user_id, {}))
    selected = user_selected_account.get(user_id, "Не выбран")
    
    text = (
        em('settings') + " <b>Менеджер аккаунтов</b>\n"
        + em('profile') + " Активных аккаунтов: " + str(count) + " (безлимит)\n"
        + em('user_check') + " Выбран: <code>" + selected + "</code>"
    )
    await safe_edit(callback.message, text, reply_markup=get_accounts_menu_keyboard())
    await callback.answer()

# ============ ЗАГРУЗКА ДИАЛОГОВ ============

async def load_dialogs(user_id: int) -> bool:
    """Загружает диалоги для выбранного аккаунта в кеш"""
    account = get_active_account(user_id)
    if not account:
        return False
    
    phone, client = account
    
    try:
        dialogs = await client.get_dialogs(limit=200)
        chats = []
        
        for dialog in dialogs:
            entity = dialog.entity
            if isinstance(entity, User):
                name = f"{entity.first_name or ''} {entity.last_name or ''}".strip()
                if not name:
                    name = f"User {entity.id}"
                chat_id = f"user_{entity.id}"
            elif isinstance(entity, (Chat, Channel)):
                name = entity.title or f"Chat {entity.id}"
                chat_id = f"chat_{entity.id}"
            else:
                continue
            chats.append((chat_id, name))
        
        if user_id not in dialogs_cache:
            dialogs_cache[user_id] = {}
        dialogs_cache[user_id][phone] = chats
        
        logger.info(f"Loaded {len(chats)} dialogs for user {user_id}, phone {phone}")
        return True
    except Exception as ex:
        logger.error(f"Error loading dialogs for user {user_id}: {ex}")
        return False

# ============ ПРОВЕРКА СПАМ БЛОКА ============

@dp.callback_query(F.data == "check_spam_block")
async def check_spam_block(callback: types.CallbackQuery):
    """Проверка спам-блока через @spambot"""
    user_id = callback.from_user.id
    account = get_active_account(user_id)
    
    if not account:
        await callback.answer("Сначала добавьте аккаунт!", show_alert=True)
        return
    
    phone, client = account
    
    status_msg = await safe_send(
        callback.message.chat.id,
        em('loading') + " <b>Проверяю спам блок...</b>"
    )
    
    try:
        # Получаем сущность бота
        spambot_entity = await client.get_entity(SPAM_BOT_USERNAME)
        
        # Отправляем /start
        await client.send_message(spambot_entity, "/start")
        
        # Ждем 4 секунды для получения ответа
        await asyncio.sleep(4)
        
        # Получаем сообщения от бота
        messages = await client.get_messages(spambot_entity, limit=10)
        
        # Ищем ответ бота (не наше сообщение)
        spam_result = "Не удалось получить ответ от @spambot"
        for msg in messages:
            if msg.text and msg.out == False:
                spam_result = msg.text
                break
        
        # Формируем результат
        result_text = (
            em('eye') + " <b>Результат проверки спам блока</b>\n\n"
            "<blockquote>" + spam_result[:1500] + "</blockquote>\n\n"
            + em('info') + " Аккаунт: <code>" + phone + "</code>"
        )
        
        # Очищаем диалог с ботом
        try:
            await client.delete_dialog(spambot_entity)
        except:
            pass
        
        await safe_edit(status_msg, result_text, reply_markup=get_functions_keyboard())
        
    except Exception as ex:
        error_text = (
            em('cross') + " <b>Ошибка при проверке спам блока</b>\n\n"
            "<code>" + str(ex) + "</code>"
        )
        try:
            await safe_edit(status_msg, error_text, reply_markup=get_functions_keyboard())
        except:
            pass
    
    await callback.answer()

# ============ РАССЫЛКА ============

@dp.callback_query(F.data == "broadcast")
async def broadcast_menu(callback: types.CallbackQuery):
    """Меню рассылки"""
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Новая рассылка",
        callback_data="new_broadcast",
        style="primary",
        icon_custom_emoji_id=eid("megaphone")
    )
    builder.button(
        text="Активные рассылки",
        callback_data="active_broadcasts",
        style="default",
        icon_custom_emoji_id=eid("clock")
    )
    builder.button(
        text="Назад",
        callback_data="main_menu",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(1)
    await safe_edit(callback.message, em('megaphone') + " <b>Рассылка</b>", reply_markup=builder.as_markup())
    await callback.answer()

@dp.callback_query(F.data == "new_broadcast")
async def new_broadcast(callback: types.CallbackQuery, state: FSMContext):
    """Начало создания новой рассылки"""
    user_id = callback.from_user.id
    account = get_active_account(user_id)
    
    if not account:
        await callback.answer("Сначала добавьте аккаунт!", show_alert=True)
        return
    
    await state.update_data(messages_list=[])
    
    text = (
        em('write') + " <b>Отправьте 1 сообщение для рассылки:</b>\n\n"
        + em('info') + " Поддерживается HTML и премиум эмодзи\n\n"
        + em('file') + " Можно добавить до 5 сообщений"
    )
    await safe_edit(callback.message, text, reply_markup=get_back_keyboard("broadcast"))
    await state.set_state(BroadcastStates.adding_messages)
    await callback.answer()

@dp.message(BroadcastStates.adding_messages)
async def broadcast_add_message(message: types.Message, state: FSMContext):
    """Добавление сообщения для рассылки"""
    msg_html = message.html_text if message.html_text else message.text or ""
    data = await state.get_data()
    messages_list = list(data.get("messages_list", []))
    
    messages_list.append(msg_html)
    await state.update_data(messages_list=messages_list)
    
    count = len(messages_list)
    
    if count >= 5:
        await safe_send(
            message.chat.id,
            em('check') + " <b>Добавлено 5/5 сообщений!</b>\n\n"
            + em('write') + " Достигнут лимит. Настройте рассылку:",
            reply_markup=get_message_actions_keyboard()
        )
    else:
        text = (
            em('check') + " <b>Сообщение " + str(count) + "/5 добавлено!</b>\n\n"
            + em('write') + " Отправьте ещё сообщение или настройте рассылку:"
        )
        await safe_send(message.chat.id, text, reply_markup=get_message_actions_keyboard())

@dp.callback_query(F.data == "add_msg", BroadcastStates.adding_messages)
async def add_more_messages(callback: types.CallbackQuery, state: FSMContext):
    """Добавление ещё одного сообщения"""
    data = await state.get_data()
    messages_list = data.get("messages_list", [])
    count = len(messages_list)
    
    if count >= 5:
        await callback.answer("Достигнут лимит в 5 сообщений!", show_alert=True)
        return
    
    text = (
        em('write') + " <b>Отправьте сообщение " + str(count + 1) + "/5:</b>\n\n"
        + em('info') + " Поддерживается HTML и премиум эмодзи"
    )
    await safe_edit(callback.message, text, reply_markup=get_back_keyboard("broadcast"))
    await callback.answer()

@dp.callback_query(F.data == "start_msg_config", BroadcastStates.adding_messages)
async def start_message_config(callback: types.CallbackQuery, state: FSMContext):
    """Переход к настройке количества сообщений"""
    data = await state.get_data()
    messages_list = data.get("messages_list", [])
    
    if not messages_list:
        await callback.answer("Добавьте хотя бы одно сообщение!", show_alert=True)
        return
    
    text = (
        em('write') + " <b>Введите количество сообщений в каждый чат:</b>\n\n"
        + em('file') + " Сообщений для рассылки: " + str(len(messages_list))
    )
    await safe_edit(callback.message, text, reply_markup=get_back_keyboard("broadcast"))
    await state.set_state(BroadcastStates.waiting_for_count)
    await callback.answer()

@dp.message(BroadcastStates.waiting_for_count)
async def broadcast_get_count(message: types.Message, state: FSMContext):
    """Ввод количества сообщений"""
    try:
        count = int(message.text)
        if count < 1:
            raise ValueError
    except ValueError:
        await safe_send(message.chat.id, em('cross') + " Введите целое положительное число")
        return
    
    await state.update_data(message_count=count)
    await safe_send(
        message.chat.id,
        em('write') + " <b>Введите задержку между сообщениями (секунд):</b>",
        reply_markup=get_back_keyboard("broadcast")
    )
    await state.set_state(BroadcastStates.waiting_for_delay)

@dp.message(BroadcastStates.waiting_for_delay)
async def broadcast_get_delay(message: types.Message, state: FSMContext):
    """Ввод задержки и переход к выбору чатов"""
    try:
        delay = float(message.text)
        if delay < 0:
            raise ValueError
    except ValueError:
        await safe_send(message.chat.id, em('cross') + " Введите положительное число")
        return
    
    await state.update_data(delay=delay, selected_chats=[], current_page=0)
    await state.set_state(BroadcastStates.selecting_chats)
    
    user_id = message.from_user.id
    status_msg = await safe_send(message.chat.id, em('loading') + " <b>Загружаем чаты...</b>")
    success = await load_dialogs(user_id)
    
    try:
        await status_msg.delete()
    except:
        pass
    
    phone = user_selected_account.get(user_id) or (
        list(user_sessions[user_id].keys())[0] if user_sessions.get(user_id) else None
    )
    
    if not success or not phone or not dialogs_cache.get(user_id, {}).get(phone):
        await safe_send(
            message.chat.id,
            em('cross') + " Не удалось загрузить чаты. Проверьте подключение аккаунта.",
            reply_markup=get_functions_keyboard()
        )
        await state.clear()
        return
    
    await delete_user_chat_messages(user_id)
    await create_chat_selection_message(message.chat.id, state, 0)

async def create_chat_selection_message(chat_id: int, state: FSMContext, page: int):
    """Создает сообщение со списком чатов для выбора"""
    user_id = chat_id
    phone = user_selected_account.get(user_id) or (
        list(user_sessions[user_id].keys())[0] if user_sessions.get(user_id) else None
    )
    if not phone:
        return
    
    chats = dialogs_cache.get(user_id, {}).get(phone, [])
    if not chats:
        await safe_send(
            chat_id,
            em('cross') + " <b>Нет доступных чатов</b>",
            reply_markup=get_functions_keyboard()
        )
        await state.clear()
        return
    
    per_page = 10
    total_pages = max(1, (len(chats) + per_page - 1) // per_page)
    start = page * per_page
    end = start + per_page
    page_chats = chats[start:end]
    
    data = await state.get_data()
    selected: List[str] = list(data.get("selected_chats", []))
    
    builder = InlineKeyboardBuilder()
    
    # Кнопки чатов
    for chat_id_str, chat_name in page_chats:
        is_selected = chat_id_str in selected
        prefix = "✅ " if is_selected else ""
        display = chat_name[:35]
        builder.button(
            text=prefix + display,
            callback_data="sc_" + chat_id_str + "_" + str(page),
            style="success" if is_selected else "default"
        )
    
    # Навигация
    nav = []
    if page > 0:
        nav.append(types.InlineKeyboardButton(
            text="◀ Назад",
            callback_data="pg_" + str(page - 1)
        ))
    if page < total_pages - 1:
        nav.append(types.InlineKeyboardButton(
            text="Вперед ▶",
            callback_data="pg_" + str(page + 1)
        ))
    if nav:
        builder.row(*nav)
    
    # Кнопки действий
    action_row = []
    if len(selected) >= 1:
        action_row.append(types.InlineKeyboardButton(
            text="🚀 Запустить (" + str(len(selected)) + ")",
            callback_data="choose_mode"
        ))
    if len(selected) > 0:
        action_row.append(types.InlineKeyboardButton(
            text="🗑 Сбросить",
            callback_data="clear_chats"
        ))
    if action_row:
        builder.row(*action_row)
    
    builder.button(
        text="Отмена",
        callback_data="cancel_broadcast",
        style="default",
        icon_custom_emoji_id=eid("cross")
    )
    builder.adjust(1)
    
    await state.update_data(current_page=page)
    
    limit_text = "до 10" if len(selected) < 10 else "ЛИМИТ"
    text = (
        em('people') + " <b>Выберите чаты для рассылки (" + limit_text + "):</b>\n"
        "Выбрано: " + str(len(selected)) + "/10 | Страница " + str(page + 1) + "/" + str(total_pages) + "\n\n"
        + em('info') + " После выбора чата появится новое сообщение\n"
        + em('send') + " Можно запустить от 1 чата"
    )
    
    msg = await safe_send(chat_id, text, reply_markup=builder.as_markup())
    await add_user_chat_message(chat_id, msg.message_id)

@dp.callback_query(F.data.startswith("sc_"))
async def select_chat(callback: types.CallbackQuery, state: FSMContext):
    """Выбор/снятие чата"""
    parts = callback.data.replace("sc_", "").split("_")
    
    if len(parts) >= 3:
        chat_id_str = parts[0] + "_" + parts[1]
        page = int(parts[2])
    elif len(parts) == 2:
        chat_id_str = parts[0]
        page = int(parts[1])
    else:
        chat_id_str = parts[0]
        page = 0
    
    data = await state.get_data()
    selected: List[str] = list(data.get("selected_chats", []))
    
    if chat_id_str in selected:
        selected.remove(chat_id_str)
        await callback.answer("Чат убран (" + str(len(selected)) + "/10)")
    else:
        if len(selected) >= 10:
            await callback.answer("Максимум 10 чатов!", show_alert=True)
            return
        selected.append(chat_id_str)
        await callback.answer("Чат добавлен (" + str(len(selected)) + "/10)")
    
    await state.update_data(selected_chats=selected, current_page=page)
    
    user_id = callback.from_user.id
    await delete_user_chat_messages(user_id)
    try:
        await callback.message.delete()
    except:
        pass
    
    await create_chat_selection_message(user_id, state, page)

@dp.callback_query(F.data.startswith("pg_"))
async def chats_page(callback: types.CallbackQuery, state: FSMContext):
    """Переключение страницы чатов"""
    page = int(callback.data.replace("pg_", ""))
    await state.update_data(current_page=page)
    
    user_id = callback.from_user.id
    await delete_user_chat_messages(user_id)
    try:
        await callback.message.delete()
    except:
        pass
    
    await create_chat_selection_message(user_id, state, page)
    await callback.answer()

@dp.callback_query(F.data == "clear_chats")
async def clear_chats_selection(callback: types.CallbackQuery, state: FSMContext):
    """Сброс выбранных чатов"""
    await state.update_data(selected_chats=[])
    data = await state.get_data()
    page = data.get("current_page", 0)
    
    user_id = callback.from_user.id
    await delete_user_chat_messages(user_id)
    try:
        await callback.message.delete()
    except:
        pass
    
    await create_chat_selection_message(user_id, state, page)
    await callback.answer("Выбор сброшен!")

@dp.callback_query(F.data == "choose_mode")
async def choose_broadcast_mode(callback: types.CallbackQuery, state: FSMContext):
    """Выбор режима рассылки"""
    data = await state.get_data()
    selected_chats = data.get("selected_chats", [])
    
    if not selected_chats:
        await callback.answer("Выберите хотя бы один чат!", show_alert=True)
        return
    
    text = (
        em('send') + " <b>Выберите режим рассылки:</b>\n\n"
        + em('loading') + " <b>Одновременная</b> — отправляет во все чаты, потом ждет задержку\n"
        + em('loading') + " <b>Рандомная</b> — отправляет в случайные чаты с задержкой"
    )
    await safe_edit(callback.message, text, reply_markup=get_mode_keyboard())
    await state.set_state(BroadcastStates.selecting_mode)
    await callback.answer()

@dp.callback_query(F.data.startswith("mode_"), BroadcastStates.selecting_mode)
async def set_broadcast_mode(callback: types.CallbackQuery, state: FSMContext):
    """Установка режима и запуск рассылки"""
    mode = callback.data.replace("mode_", "")
    await state.update_data(mode=mode)
    await start_broadcast_execution(callback, state)

async def start_broadcast_execution(callback: types.CallbackQuery, state: FSMContext):
    """Запуск выполнения рассылки"""
    data = await state.get_data()
    selected_chats = data.get("selected_chats", [])
    mode = data.get("mode", "sync")
    
    if not selected_chats:
        await callback.answer("Выберите хотя бы один чат!", show_alert=True)
        return
    
    user_id = callback.from_user.id
    broadcast_id = "bc_" + str(int(datetime.now().timestamp()))
    
    if user_id not in active_broadcasts:
        active_broadcasts[user_id] = {}
    
    account = get_active_account(user_id)
    if not account:
        await callback.answer("Нет активного аккаунта!", show_alert=True)
        return
    
    phone, client = account
    
    task = asyncio.create_task(
        execute_broadcast(
            user_id=user_id,
            broadcast_id=broadcast_id,
            chats=selected_chats,
            messages_list=data.get("messages_list", [""]),
            count=data.get("message_count", 1),
            delay=data.get("delay", 1),
            client=client,
            mode=mode
        )
    )
    active_broadcasts[user_id][broadcast_id] = task
    
    await delete_user_chat_messages(user_id)
    
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Остановить",
        callback_data="stop_" + broadcast_id,
        style="danger",
        icon_custom_emoji_id=eid("cross")
    )
    builder.button(
        text="В меню",
        callback_data="main_menu",
        style="default",
        icon_custom_emoji_id=eid("home")
    )
    builder.adjust(1)
    
    mode_text = "Одновременная" if mode == "sync" else "Рандомная"
    text = (
        em('send') + " <b>Рассылка запущена!</b>\n"
        + em('people') + " Чатов: " + str(len(selected_chats)) + "\n"
        + em('file') + " Сообщений: " + str(len(data.get("messages_list", ['']))) 
        + " (в каждый: " + str(data.get('message_count', 1)) + ")\n"
        + em('clock') + " Задержка: " + str(data.get('delay', 1)) + "с\n"
        + em('loading') + " Режим: " + mode_text
    )
    
    try:
        await callback.message.delete()
    except:
        pass
    
    await safe_send(user_id, text, reply_markup=builder.as_markup())
    await state.clear()
    await callback.answer("Рассылка запущена!", show_alert=True)

async def execute_broadcast(
    user_id: int, broadcast_id: str, chats: List[str],
    messages_list: List[str], count: int, delay: float,
    client: TelegramClient, mode: str
):
    """Выполнение рассылки в фоновом режиме"""
    completed = 0
    errors = 0
    total_messages = len(messages_list)
    
    if mode == "sync":
        # Одновременная рассылка
        for i in range(count):
            if broadcast_id not in active_broadcasts.get(user_id, {}):
                break
            
            msg = random.choice(messages_list) if total_messages > 1 else messages_list[0]
            
            for chat_id_str in chats:
                if broadcast_id not in active_broadcasts.get(user_id, {}):
                    break
                
                try:
                    entity_id = int(chat_id_str.split("_")[1])
                    entity = await client.get_entity(entity_id)
                    
                    try:
                        await client.send_message(entity, msg, parse_mode='html')
                        completed += 1
                    except Exception:
                        errors += 1
                        try:
                            clean = re.sub(r'<[^>]+>', '', msg)
                            await client.send_message(entity, clean)
                            completed += 1
                        except:
                            pass
                except Exception:
                    errors += 1
            
            if i < count - 1 and broadcast_id in active_broadcasts.get(user_id, {}):
                await asyncio.sleep(delay)
    else:
        # Рандомная рассылка
        total_operations = count * len(chats)
        chat_indices = list(range(len(chats)))
        
        for i in range(total_operations):
            if broadcast_id not in active_broadcasts.get(user_id, {}):
                break
            
            chat_idx = random.choice(chat_indices)
            chat_id_str = chats[chat_idx]
            msg = random.choice(messages_list) if total_messages > 1 else messages_list[0]
            
            try:
                entity_id = int(chat_id_str.split("_")[1])
                entity = await client.get_entity(entity_id)
                
                try:
                    await client.send_message(entity, msg, parse_mode='html')
                    completed += 1
                except Exception:
                    errors += 1
                    try:
                        clean = re.sub(r'<[^>]+>', '', msg)
                        await client.send_message(entity, clean)
                        completed += 1
                    except:
                        pass
            except Exception:
                errors += 1
            
            if broadcast_id in active_broadcasts.get(user_id, {}):
                await asyncio.sleep(delay)
    
    logger.info(
        f"Broadcast {broadcast_id} done. "
        f"Mode: {mode}, OK: {completed}, ERR: {errors}"
    )
    active_broadcasts.get(user_id, {}).pop(broadcast_id, None)

@dp.callback_query(F.data == "active_broadcasts")
async def active_broadcasts_list(callback: types.CallbackQuery):
    """Список активных рассылок"""
    user_id = callback.from_user.id
    broadcasts = active_broadcasts.get(user_id, {})
    
    if not broadcasts:
        builder = InlineKeyboardBuilder()
        builder.button(
            text="Назад",
            callback_data="broadcast",
            style="default",
            icon_custom_emoji_id=eid("back")
        )
        text = em('unlock') + " <b>Нет активных рассылок</b>"
        await safe_edit(callback.message, text, reply_markup=builder.as_markup())
        await callback.answer()
        return
    
    builder = InlineKeyboardBuilder()
    for bid in broadcasts:
        builder.button(
            text="Остановить " + bid[:8] + "...",
            callback_data="stop_" + bid,
            style="danger",
            icon_custom_emoji_id=eid("cross")
        )
    builder.button(
        text="Назад",
        callback_data="broadcast",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(1)
    
    text = em('clock') + " <b>Активные рассылки:</b> " + str(len(broadcasts))
    await safe_edit(callback.message, text, reply_markup=builder.as_markup())
    await callback.answer()

@dp.callback_query(F.data.startswith("stop_"))
async def stop_broadcast(callback: types.CallbackQuery):
    """Остановка активной рассылки"""
    broadcast_id = callback.data.replace("stop_", "")
    user_id = callback.from_user.id
    
    if user_id in active_broadcasts and broadcast_id in active_broadcasts[user_id]:
        active_broadcasts[user_id][broadcast_id].cancel()
        del active_broadcasts[user_id][broadcast_id]
        await callback.answer("Рассылка остановлена!", show_alert=True)
    else:
        await callback.answer("Рассылка не найдена", show_alert=True)
    
    await active_broadcasts_list(callback)

@dp.callback_query(F.data == "cancel_broadcast")
async def cancel_broadcast(callback: types.CallbackQuery, state: FSMContext):
    """Отмена создания рассылки"""
    user_id = callback.from_user.id
    await delete_user_chat_messages(user_id)
    await state.clear()
    try:
        await callback.message.delete()
    except:
        pass
    await safe_send(
        callback.message.chat.id,
        em('cross') + " <b>Рассылка отменена</b>",
        reply_markup=get_functions_keyboard()
    )
    await callback.answer()

# ============ ВСТУПЛЕНИЕ В ЧАТЫ ============

@dp.callback_query(F.data == "join_chats")
async def join_chats_menu(callback: types.CallbackQuery, state: FSMContext):
    """Меню вступления в чаты"""
    user_id = callback.from_user.id
    account = get_active_account(user_id)
    
    if not account:
        await callback.answer("Сначала добавьте аккаунт!", show_alert=True)
        return
    
    text = (
        em('people') + " <b>Вступление в чаты</b>\n\n"
        + em('write') + " Отправьте список юзернеймов (каждый с новой строки):\n"
        "<code>@chat1\n@chat2\n@chat3</code>\n\n"
        + em('info') + " Задержка между вступлениями: 15 секунд"
    )
    await safe_edit(callback.message, text, reply_markup=get_back_keyboard("functions"))
    await state.set_state(JoinStates.waiting_for_usernames)
    await callback.answer()

@dp.message(JoinStates.waiting_for_usernames)
async def process_join_usernames(message: types.Message, state: FSMContext):
    """Обработка списка юзернеймов для вступления"""
    text = message.text.strip()
    lines = []
    
    for line in text.split('\n'):
        line = line.strip()
        if not line:
            continue
        # Очищаем от ссылок и @
        line = line.replace('https://t.me/', '')
        line = line.replace('@', '')
        line = line.strip()
        if line:
            lines.append(line)
    
    if not lines:
        await safe_send(
            message.chat.id,
            em('cross') + " Не найдено ни одного юзернейма",
            reply_markup=get_back_keyboard("functions")
        )
        return
    
    if len(lines) > 50:
        await safe_send(
            message.chat.id,
            em('cross') + " Максимум 50 чатов за раз",
            reply_markup=get_back_keyboard("functions")
        )
        return
    
    user_id = message.from_user.id
    account = get_active_account(user_id)
    
    if not account:
        await safe_send(
            message.chat.id,
            em('cross') + " Нет активных аккаунтов",
            reply_markup=get_functions_keyboard()
        )
        await state.clear()
        return
    
    phone, client = account
    await state.clear()
    
    status_msg = await safe_send(
        message.chat.id,
        em('loading') + " <b>Начинаю вступление...</b>\n"
        + em('people') + " Всего: " + str(len(lines)) + "\n"
        + em('clock') + " Задержка: 15с"
    )
    
    joined = 0
    failed = 0
    failed_list = []
    
    for i, username in enumerate(lines, 1):
        try:
            try:
                await client(JoinChannelRequest(username))
                joined += 1
            except FloodWaitError as e:
                wait_time = e.seconds
                await safe_edit(
                    status_msg,
                    em('loading') + " <b>FloodWait " + str(wait_time) + "с...</b>\n"
                    "Прогресс: " + str(i) + "/" + str(len(lines))
                )
                await asyncio.sleep(wait_time)
                try:
                    await client(JoinChannelRequest(username))
                    joined += 1
                except:
                    failed += 1
                    failed_list.append(username)
            except:
                try:
                    entity = await client.get_entity(username)
                    await client(JoinChannelRequest(entity))
                    joined += 1
                except:
                    failed += 1
                    failed_list.append(username)
        except:
            failed += 1
            failed_list.append(username)
        
        # Обновление прогресса
        if i % 3 == 0 or i == len(lines):
            progress_text = (
                em('loading') + " <b>Вступление в чаты...</b>\n"
                "Прогресс: " + str(i) + "/" + str(len(lines)) + "\n"
                + em('check') + " Вступил: " + str(joined) + "\n"
                + em('cross') + " Ошибок: " + str(failed)
            )
            try:
                await safe_edit(status_msg, progress_text)
            except:
                pass
        
        if i < len(lines):
            await asyncio.sleep(15)
    
    # Финальный отчет
    result_text = (
        em('celebration') + " <b>Вступление завершено!</b>\n\n"
        + em('check') + " Успешно: " + str(joined) + "\n"
        + em('cross') + " Ошибок: " + str(failed)
    )
    
    if failed_list:
        result_text += "\n\n" + em('cross') + " <b>Не удалось вступить:</b>\n"
        result_text += "\n".join("• " + u for u in failed_list[:10])
        if len(failed_list) > 10:
            result_text += "\n... и еще " + str(len(failed_list) - 10)
    
    await safe_edit(status_msg, result_text, reply_markup=get_functions_keyboard())

# ============ МАСС РЕАКЦИИ ============

@dp.callback_query(F.data == "mass_reactions")
async def mass_reactions_menu(callback: types.CallbackQuery):
    """Меню масс-реакций"""
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Запустить реакции",
        callback_data="start_reactions",
        style="primary",
        icon_custom_emoji_id=eid("smile")
    )
    builder.button(
        text="Активные реакции",
        callback_data="active_reactions_list",
        style="default",
        icon_custom_emoji_id=eid("clock")
    )
    builder.button(
        text="Назад",
        callback_data="main_menu",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(1)
    await safe_edit(
        callback.message,
        em('smile') + " <b>Масс реакции</b>",
        reply_markup=builder.as_markup()
    )
    await callback.answer()

@dp.callback_query(F.data == "start_reactions")
async def start_reactions(callback: types.CallbackQuery, state: FSMContext):
    """Начало настройки масс-реакций"""
    user_id = callback.from_user.id
    account = get_active_account(user_id)
    
    if not account:
        await callback.answer("Сначала добавьте аккаунт!", show_alert=True)
        return
    
    await state.clear()
    await state.set_state(ReactionStates.selecting_chats)
    await state.update_data(selected_chats=[], current_page=0)
    
    status_msg = await safe_send(
        callback.message.chat.id,
        em('loading') + " <b>Загружаем чаты...</b>"
    )
    success = await load_dialogs(user_id)
    
    try:
        await status_msg.delete()
    except:
        pass
    
    account_result = get_active_account(user_id)
    if not account_result:
        await safe_send(
            callback.message.chat.id,
            em('cross') + " Нет активного аккаунта",
            reply_markup=get_functions_keyboard()
        )
        await state.clear()
        return
    
    phone = account_result[0]
    if not success or not dialogs_cache.get(user_id, {}).get(phone):
        await safe_send(
            callback.message.chat.id,
            em('cross') + " Не удалось загрузить чаты.",
            reply_markup=get_functions_keyboard()
        )
        await state.clear()
        return
    
    await delete_user_chat_messages(user_id)
    await create_reaction_chat_selection(callback.message.chat.id, state, 0)
    await callback.answer()

async def create_reaction_chat_selection(chat_id: int, state: FSMContext, page: int):
    """Создает сообщение со списком чатов для выбора реакций"""
    user_id = chat_id
    account_result = get_active_account(user_id)
    if not account_result:
        return
    
    phone = account_result[0]
    chats = dialogs_cache.get(user_id, {}).get(phone, [])
    
    if not chats:
        await safe_send(
            chat_id,
            em('cross') + " <b>Нет доступных чатов</b>",
            reply_markup=get_functions_keyboard()
        )
        await state.clear()
        return
    
    per_page = 10
    total_pages = max(1, (len(chats) + per_page - 1) // per_page)
    start = page * per_page
    end = start + per_page
    page_chats = chats[start:end]
    
    data = await state.get_data()
    selected: List[str] = list(data.get("selected_chats", []))
    
    builder = InlineKeyboardBuilder()
    
    for chat_id_str, chat_name in page_chats:
        is_selected = chat_id_str in selected
        prefix = "✅ " if is_selected else ""
        display = chat_name[:35]
        builder.button(
            text=prefix + display,
            callback_data="rsc_" + chat_id_str + "_" + str(page),
            style="success" if is_selected else "default"
        )
    
    nav = []
    if page > 0:
        nav.append(types.InlineKeyboardButton(
            text="◀ Назад", callback_data="rpg_" + str(page - 1)
        ))
    if page < total_pages - 1:
        nav.append(types.InlineKeyboardButton(
            text="Вперед ▶", callback_data="rpg_" + str(page + 1)
        ))
    if nav:
        builder.row(*nav)
    
    action_row = []
    if len(selected) >= 1:
        action_row.append(types.InlineKeyboardButton(
            text="🚀 Далее (" + str(len(selected)) + ")",
            callback_data="react_next"
        ))
    if len(selected) > 0:
        action_row.append(types.InlineKeyboardButton(
            text="🗑 Сбросить",
            callback_data="rclear_chats"
        ))
    if action_row:
        builder.row(*action_row)
    
    builder.button(
        text="Отмена",
        callback_data="cancel_reactions",
        style="default",
        icon_custom_emoji_id=eid("cross")
    )
    builder.adjust(1)
    
    await state.update_data(current_page=page)
    
    limit_text = "до 5" if len(selected) < 5 else "ЛИМИТ"
    text = (
        em('people') + " <b>Выберите чаты для реакций (" + limit_text + "):</b>\n"
        "Выбрано: " + str(len(selected)) + "/5 | Страница " + str(page + 1) + "/" + str(total_pages)
    )
    
    msg = await safe_send(chat_id, text, reply_markup=builder.as_markup())
    await add_user_chat_message(chat_id, msg.message_id)

@dp.callback_query(F.data.startswith("rsc_"))
async def reaction_select_chat(callback: types.CallbackQuery, state: FSMContext):
    """Выбор чата для реакций"""
    parts = callback.data.replace("rsc_", "").split("_")
    
    if len(parts) >= 3:
        chat_id_str = parts[0] + "_" + parts[1]
        page = int(parts[2])
    elif len(parts) == 2:
        chat_id_str = parts[0]
        page = int(parts[1])
    else:
        chat_id_str = parts[0]
        page = 0
    
    data = await state.get_data()
    selected: List[str] = list(data.get("selected_chats", []))
    
    if chat_id_str in selected:
        selected.remove(chat_id_str)
        await callback.answer("Чат убран (" + str(len(selected)) + "/5)")
    else:
        if len(selected) >= 5:
            await callback.answer("Максимум 5 чатов!", show_alert=True)
            return
        selected.append(chat_id_str)
        await callback.answer("Чат добавлен (" + str(len(selected)) + "/5)")
    
    await state.update_data(selected_chats=selected)
    
    user_id = callback.from_user.id
    await delete_user_chat_messages(user_id)
    try:
        await callback.message.delete()
    except:
        pass
    
    await create_reaction_chat_selection(user_id, state, page)

@dp.callback_query(F.data.startswith("rpg_"))
async def reaction_chats_page(callback: types.CallbackQuery, state: FSMContext):
    """Переключение страницы чатов для реакций"""
    page = int(callback.data.replace("rpg_", ""))
    await state.update_data(current_page=page)
    
    user_id = callback.from_user.id
    await delete_user_chat_messages(user_id)
    try:
        await callback.message.delete()
    except:
        pass
    
    await create_reaction_chat_selection(user_id, state, page)
    await callback.answer()

@dp.callback_query(F.data == "rclear_chats")
async def reaction_clear_chats(callback: types.CallbackQuery, state: FSMContext):
    """Сброс выбранных чатов для реакций"""
    await state.update_data(selected_chats=[])
    data = await state.get_data()
    page = data.get("current_page", 0)
    
    user_id = callback.from_user.id
    await delete_user_chat_messages(user_id)
    try:
        await callback.message.delete()
    except:
        pass
    
    await create_reaction_chat_selection(user_id, state, page)
    await callback.answer("Выбор сброшен!")

@dp.callback_query(F.data == "react_next")
async def reaction_next_step(callback: types.CallbackQuery, state: FSMContext):
    """Переход к вводу задержки для реакций"""
    data = await state.get_data()
    selected_chats = data.get("selected_chats", [])
    
    if not selected_chats:
        await callback.answer("Выберите хотя бы один чат!", show_alert=True)
        return
    
    user_id = callback.from_user.id
    await delete_user_chat_messages(user_id)
    
    try:
        await callback.message.delete()
    except:
        pass
    
    text = em('write') + " <b>Введите задержку между реакциями (секунд):</b>"
    await safe_send(user_id, text, reply_markup=get_back_keyboard("cancel_reactions"))
    await state.set_state(ReactionStates.waiting_for_delay)
    await callback.answer()

@dp.message(ReactionStates.waiting_for_delay)
async def reaction_get_delay(message: types.Message, state: FSMContext):
    """Ввод задержки для реакций"""
    try:
        delay = float(message.text)
        if delay < 1:
            raise ValueError
    except ValueError:
        await safe_send(message.chat.id, em('cross') + " Введите положительное число (минимум 1 секунда)")
        return
    
    await state.update_data(delay=delay)
    await state.set_state(ReactionStates.selecting_type)
    
    text = (
        em('send') + " <b>Выберите тип работы:</b>\n\n"
        + em('loading') + " <b>Всё подряд</b> — ставит реакцию на все новые сообщения\n"
        + em('loading') + " <b>Через один</b> — ставит реакцию, пропуская одно сообщение"
    )
    await safe_send(message.chat.id, text, reply_markup=get_reaction_type_keyboard())

@dp.callback_query(F.data.startswith("rtype_"), ReactionStates.selecting_type)
async def reaction_set_type(callback: types.CallbackQuery, state: FSMContext):
    """Выбор типа реакций"""
    rtype = callback.data.replace("rtype_", "")
    await state.update_data(reaction_type=rtype)
    await state.set_state(ReactionStates.selecting_reaction)
    
    await safe_edit(
        callback.message,
        em('smile') + " <b>Выберите реакцию:</b>",
        reply_markup=get_reactions_keyboard()
    )
    await callback.answer()

@dp.callback_query(F.data.startswith("react_"), ReactionStates.selecting_reaction)
async def reaction_start(callback: types.CallbackQuery, state: FSMContext):
    """Запуск масс-реакций"""
    reaction = callback.data.replace("react_", "")
    data = await state.get_data()
    selected_chats = data.get("selected_chats", [])
    delay = data.get("delay", 5)
    rtype = data.get("reaction_type", "all")
    
    if not selected_chats:
        await callback.answer("Нет выбранных чатов!", show_alert=True)
        return
    
    user_id = callback.from_user.id
    account = get_active_account(user_id)
    if not account:
        await callback.answer("Нет активного аккаунта!", show_alert=True)
        return
    
    phone, client = account
    reaction_id = "reaction_" + str(int(datetime.now().timestamp()))
    
    if user_id not in active_reactions:
        active_reactions[user_id] = {}
    
    task = asyncio.create_task(
        run_reactions(
            user_id=user_id,
            reaction_id=reaction_id,
            chats=selected_chats,
            delay=delay,
            rtype=rtype,
            reaction=reaction,
            client=client
        )
    )
    
    active_reactions[user_id][reaction_id] = {
        "task": task,
        "chats": selected_chats,
        "reaction": reaction,
        "type": rtype,
        "delay": delay
    }
    
    text = (
        em('smile') + " <b>Масс реакции запущены!</b>\n"
        + em('people') + " Чатов: " + str(len(selected_chats)) + "\n"
        + em('clock') + " Задержка: " + str(delay) + "с\n"
        + em('send') + " Реакция: " + reaction + "\n"
        + em('loading') + " Тип: " + ("Всё подряд" if rtype == "all" else "Через один")
    )
    
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Остановить",
        callback_data="rstop_" + reaction_id,
        style="danger",
        icon_custom_emoji_id=eid("cross")
    )
    builder.button(
        text="В меню",
        callback_data="main_menu",
        style="default",
        icon_custom_emoji_id=eid("home")
    )
    builder.adjust(1)
    
    try:
        await callback.message.delete()
    except:
        pass
    
    await safe_send(callback.message.chat.id, text, reply_markup=builder.as_markup())
    await state.clear()
    await callback.answer("Реакции запущены!", show_alert=True)

async def run_reactions(
    user_id: int, reaction_id: str, chats: List[str],
    delay: float, rtype: str, reaction: str, client: TelegramClient
):
    """Запуск отслеживания новых сообщений и установки реакций"""
    message_counters = {}
    chat_ids = []
    
    for chat_str in chats:
        try:
            chat_id = int(chat_str.split("_")[1])
            chat_ids.append(chat_id)
        except:
            pass
    
    @client.on(events.NewMessage(chats=chat_ids))
    async def reaction_handler(event):
        if reaction_id not in active_reactions.get(user_id, {}):
            return
        
        chat_id = event.chat_id
        if chat_id not in message_counters:
            message_counters[chat_id] = 0
        
        message_counters[chat_id] += 1
        
        if rtype == "skip":
            if message_counters[chat_id] % 2 == 0:
                return
        
        try:
            await asyncio.sleep(delay)
            await client(SendReactionRequest(
                peer=event.message.peer_id,
                msg_id=event.message.id,
                reaction=[ReactionEmoji(emoticon=reaction)]
            ))
        except Exception as ex:
            logger.error(f"Reaction error: {ex}")
    
    # Держим задачу активной пока реакции не остановлены
    while reaction_id in active_reactions.get(user_id, {}):
        await asyncio.sleep(1)
    
    client.remove_event_handler(reaction_handler)
    logger.info(f"Reactions stopped for {reaction_id}")

@dp.callback_query(F.data == "active_reactions_list")
async def active_reactions_list(callback: types.CallbackQuery):
    """Список активных реакций"""
    user_id = callback.from_user.id
    reactions = active_reactions.get(user_id, {})
    
    if not reactions:
        builder = InlineKeyboardBuilder()
        builder.button(
            text="Назад",
            callback_data="mass_reactions",
            style="default",
            icon_custom_emoji_id=eid("back")
        )
        text = em('unlock') + " <b>Нет активных реакций</b>"
        await safe_edit(callback.message, text, reply_markup=builder.as_markup())
        await callback.answer()
        return
    
    builder = InlineKeyboardBuilder()
    for rid, rdata in reactions.items():
        builder.button(
            text="Остановить " + rdata["reaction"] + " (" + str(len(rdata["chats"])) + " чатов)",
            callback_data="rstop_" + rid,
            style="danger",
            icon_custom_emoji_id=eid("cross")
        )
    builder.button(
        text="Назад",
        callback_data="mass_reactions",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(1)
    
    text = em('smile') + " <b>Активные реакции:</b> " + str(len(reactions))
    await safe_edit(callback.message, text, reply_markup=builder.as_markup())
    await callback.answer()

@dp.callback_query(F.data.startswith("rstop_"))
async def stop_reactions(callback: types.CallbackQuery):
    """Остановка активных реакций"""
    reaction_id = callback.data.replace("rstop_", "")
    user_id = callback.from_user.id
    
    if user_id in active_reactions and reaction_id in active_reactions[user_id]:
        active_reactions[user_id][reaction_id]["task"].cancel()
        del active_reactions[user_id][reaction_id]
        await callback.answer("Реакции остановлены!", show_alert=True)
    else:
        await callback.answer("Реакции не найдены", show_alert=True)
    
    await active_reactions_list(callback)

@dp.callback_query(F.data == "cancel_reactions")
async def cancel_reactions(callback: types.CallbackQuery, state: FSMContext):
    """Отмена настройки реакций"""
    user_id = callback.from_user.id
    await delete_user_chat_messages(user_id)
    await state.clear()
    try:
        await callback.message.delete()
    except:
        pass
    await safe_send(
        callback.message.chat.id,
        em('cross') + " <b>Реакции отменены</b>",
        reply_markup=get_functions_keyboard()
    )
    await callback.answer()

# ============ СИСТЕМА КОМАНД ============

@dp.callback_query(F.data == "commands_menu")
async def commands_menu(callback: types.CallbackQuery):
    """Меню команд"""
    user_id = callback.from_user.id
    cmds = user_commands.get(user_id, {})
    real_cmds = {k: v for k, v in cmds.items() if not k.startswith("__")}
    welcome = user_welcome_enabled.get(user_id, False)
    configs = saved_configs.get(user_id, {})
    
    text = (
        em('code') + " <b>Команды</b>\n\n"
        + em('file') + " Команд: " + str(len(real_cmds)) + "/20\n"
        + em('gift') + " Приветствие: " + ("Вкл" if welcome else "Выкл") + "\n"
        + em('box') + " Конфигов: " + str(len(configs)) + "\n\n"
        + em('info') + " Команды начинаются с точки (.)\n"
        + em('info') + " Плейсхолдеры: {NICK}, {FIRSTNAME}, {LASTNAME}, {FULLNAME}, {USERNAME}, {ID}\n"
        + em('info') + " Команды отвечают всегда, даже в существующих диалогах\n"
        + em('info') + " Поддерживается HTML и премиум эмодзи"
    )
    await safe_edit(callback.message, text, reply_markup=get_commands_keyboard())
    await callback.answer()

@dp.callback_query(F.data == "add_command")
async def add_command(callback: types.CallbackQuery, state: FSMContext):
    """Добавление новой команды"""
    user_id = callback.from_user.id
    cmds = user_commands.get(user_id, {})
    real_cmds = {k: v for k, v in cmds.items() if not k.startswith("__")}
    
    if len(real_cmds) >= 20:
        await callback.answer("Достигнут лимит в 20 команд!", show_alert=True)
        return
    
    text = (
        em('write') + " <b>Добавление команды</b>\n\n"
        + em('info') + " Отправьте название команды (с точкой):\n"
        "<code>.прайс</code>\n<code>.хелп</code>\n<code>.контакты</code>"
    )
    await safe_edit(callback.message, text, reply_markup=get_back_keyboard("commands_menu"))
    await state.set_state(CommandStates.waiting_for_command_name)
    await callback.answer()

@dp.message(CommandStates.waiting_for_command_name)
async def command_get_name(message: types.Message, state: FSMContext):
    """Получение названия команды"""
    cmd_name = message.text.strip().lower()
    
    if not cmd_name.startswith('.'):
        await safe_send(message.chat.id, em('cross') + " Команда должна начинаться с точки (.)")
        return
    
    if len(cmd_name) < 2:
        await safe_send(message.chat.id, em('cross') + " Слишком короткая команда")
        return
    
    user_id = message.from_user.id
    cmds = user_commands.get(user_id, {})
    
    if cmd_name in cmds:
        await safe_send(message.chat.id, em('cross') + " Такая команда уже существует!")
        return
    
    await state.update_data(command_name=cmd_name)
    
    text = (
        em('write') + " <b>Команда: <code>" + cmd_name + "</code></b>\n\n"
        + em('info') + " Теперь отправьте ответ на эту команду:\n"
        + em('info') + " Можно использовать плейсхолдеры:\n"
        "<code>{NICK}</code> - имя пользователя\n"
        "<code>{FIRSTNAME}</code> - имя\n"
        "<code>{LASTNAME}</code> - фамилия\n"
        "<code>{FULLNAME}</code> - полное имя\n"
        "<code>{USERNAME}</code> - юзернейм\n"
        "<code>{ID}</code> - ID пользователя\n\n"
        + em('info') + " Поддерживается HTML и премиум эмодзи"
    )
    await safe_send(message.chat.id, text, reply_markup=get_back_keyboard("commands_menu"))
    await state.set_state(CommandStates.waiting_for_command_response)

@dp.message(CommandStates.waiting_for_command_response)
async def command_get_response(message: types.Message, state: FSMContext):
    """Получение ответа на команду"""
    response = message.html_text if message.html_text else message.text or ""
    data = await state.get_data()
    cmd_name = data.get("command_name", ".unknown")
    
    user_id = message.from_user.id
    if user_id not in user_commands:
        user_commands[user_id] = {}
    
    user_commands[user_id][cmd_name] = response
    await state.clear()
    
    # Запускаем обработчик команд
    await setup_command_handler(user_id)
    
    text = (
        em('check') + " <b>Команда добавлена!</b>\n"
        "<code>" + cmd_name + "</code>"
    )
    await safe_send(message.chat.id, text, reply_markup=get_commands_keyboard())

async def setup_command_handler(user_id: int):
    """Настраивает обработчик входящих сообщений для команд"""
    account = get_active_account(user_id)
    if not account:
        return
    
    phone, client = account
    
    # Удаляем старый обработчик если есть
    if user_id in active_auto_reply_handlers:
        try:
            old_handler = active_auto_reply_handlers[user_id].get("cmd_handler")
            if old_handler:
                client.remove_event_handler(old_handler)
        except:
            pass
    
    @client.on(events.NewMessage(incoming=True))
    async def command_handler(event):
        if not event.is_private:
            return
        
        msg_text = event.message.text or ""
        if not msg_text:
            return
        
        msg_text = msg_text.strip().lower()
        cmds = user_commands.get(user_id, {})
        
        if msg_text in cmds:
            try:
                # Получаем информацию об отправителе
                sender = await event.get_sender()
                user_info = {
                    "first_name": getattr(sender, 'first_name', '') or "",
                    "last_name": getattr(sender, 'last_name', '') or "",
                    "full_name": "",
                    "username": getattr(sender, 'username', '') or "",
                    "id": str(event.sender_id),
                    "phone": getattr(sender, 'phone', '') or "",
                }
                user_info["full_name"] = f"{user_info['first_name']} {user_info['last_name']}".strip()
                
                # Обрабатываем плейсхолдеры
                response = process_placeholders(cmds[msg_text], user_info)
                
                await client.send_message(
                    event.sender_id,
                    response,
                    parse_mode='html'
                )
            except Exception as ex:
                logger.error(f"Command reply error: {ex}")
    
    if user_id not in active_auto_reply_handlers:
        active_auto_reply_handlers[user_id] = {}
    active_auto_reply_handlers[user_id]["cmd_handler"] = command_handler
    active_auto_reply_handlers[user_id]["client"] = client
    
    # Также настраиваем приветствие если включено
    if user_welcome_enabled.get(user_id, False):
        await setup_welcome_handler(user_id)

@dp.callback_query(F.data == "welcome_setup")
async def welcome_setup(callback: types.CallbackQuery, state: FSMContext):
    """Настройка приветственного сообщения"""
    user_id = callback.from_user.id
    welcome_enabled = user_welcome_enabled.get(user_id, False)
    
    if welcome_enabled:
        # Выключаем приветствие
        user_welcome_enabled[user_id] = False
        # Удаляем обработчик приветствия
        if user_id in active_auto_reply_handlers:
            try:
                old_handler = active_auto_reply_handlers[user_id].get("welcome_handler")
                client = active_auto_reply_handlers[user_id].get("client")
                if old_handler and client:
                    client.remove_event_handler(old_handler)
            except:
                pass
        await callback.answer("Приветствие отключено!", show_alert=True)
        await commands_menu(callback)
        return
    
    text = (
        em('gift') + " <b>Настройка приветствия</b>\n\n"
        + em('write') + " Отправьте сообщение-приветствие для новых пользователей:\n"
        + em('info') + " Поддерживается HTML и премиум эмодзи\n"
        + em('info') + " Можно использовать: {NICK}, {FIRSTNAME}, {LASTNAME}, {FULLNAME}\n"
        + em('info') + " Бот будет отвечать на ПЕРВОЕ сообщение от нового пользователя (пустой чат)"
    )
    await safe_edit(callback.message, text, reply_markup=get_back_keyboard("commands_menu"))
    await state.set_state(WelcomeStates.waiting_for_welcome_message)
    await callback.answer()

@dp.message(WelcomeStates.waiting_for_welcome_message)
async def welcome_get_message(message: types.Message, state: FSMContext):
    """Получение приветственного сообщения"""
    welcome_msg = message.html_text if message.html_text else message.text or ""
    user_id = message.from_user.id
    
    if user_id not in user_commands:
        user_commands[user_id] = {}
    
    user_commands[user_id]["__welcome__"] = welcome_msg
    user_welcome_enabled[user_id] = True
    await state.clear()
    
    await setup_welcome_handler(user_id)
    
    text = em('check') + " <b>Приветствие настроено!</b>"
    await safe_send(message.chat.id, text, reply_markup=get_commands_keyboard())

async def setup_welcome_handler(user_id: int):
    """Настраивает обработчик приветствия для новых пользователей"""
    account = get_active_account(user_id)
    if not account:
        return
    
    phone, client = account
    welcomed_users = set()
    
    # Удаляем старый обработчик если есть
    if user_id in active_auto_reply_handlers:
        try:
            old_handler = active_auto_reply_handlers[user_id].get("welcome_handler")
            if old_handler:
                client.remove_event_handler(old_handler)
        except:
            pass
    
    @client.on(events.NewMessage(incoming=True))
    async def welcome_handler(event):
        if not event.is_private:
            return
        
        sender_id = event.sender_id
        if sender_id in welcomed_users:
            return
        
        # Проверяем что это первое сообщение от пользователя
        try:
            messages = await client.get_messages(sender_id, limit=2)
            if len(messages) <= 1:
                welcomed_users.add(sender_id)
                welcome_msg = user_commands.get(user_id, {}).get("__welcome__", "")
                if welcome_msg:
                    sender = await event.get_sender()
                    user_info = {
                        "first_name": getattr(sender, 'first_name', '') or "",
                        "last_name": getattr(sender, 'last_name', '') or "",
                        "full_name": "",
                        "username": getattr(sender, 'username', '') or "",
                        "id": str(event.sender_id),
                        "phone": "",
                    }
                    user_info["full_name"] = f"{user_info['first_name']} {user_info['last_name']}".strip()
                    response = process_placeholders(welcome_msg, user_info)
                    try:
                        await client.send_message(sender_id, response, parse_mode='html')
                    except Exception as ex:
                        logger.error(f"Welcome message error: {ex}")
        except:
            pass
    
    if user_id not in active_auto_reply_handlers:
        active_auto_reply_handlers[user_id] = {}
    active_auto_reply_handlers[user_id]["welcome_handler"] = welcome_handler
    if "client" not in active_auto_reply_handlers[user_id]:
        active_auto_reply_handlers[user_id]["client"] = client

@dp.callback_query(F.data == "my_commands")
async def my_commands_list(callback: types.CallbackQuery):
    """Список команд пользователя"""
    user_id = callback.from_user.id
    cmds = user_commands.get(user_id, {})
    display_cmds = {k: v for k, v in cmds.items() if not k.startswith("__")}
    
    if not display_cmds:
        text = em('unlock') + " <b>Нет добавленных команд</b>"
        builder = InlineKeyboardBuilder()
        builder.button(
            text="Назад",
            callback_data="commands_menu",
            style="default",
            icon_custom_emoji_id=eid("back")
        )
        await safe_edit(callback.message, text, reply_markup=builder.as_markup())
        await callback.answer()
        return
    
    builder = InlineKeyboardBuilder()
    for cmd_name in display_cmds:
        builder.button(text=cmd_name, callback_data="edit_cmd_" + cmd_name, style="default")
    builder.button(
        text="Назад",
        callback_data="commands_menu",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(2)
    
    text = em('file') + " <b>Мои команды (" + str(len(display_cmds)) + "/20):</b>"
    await safe_edit(callback.message, text, reply_markup=builder.as_markup())
    await callback.answer()

@dp.callback_query(F.data.startswith("edit_cmd_"))
async def edit_command(callback: types.CallbackQuery):
    """Просмотр и управление командой"""
    cmd_name = callback.data.replace("edit_cmd_", "")
    user_id = callback.from_user.id
    cmds = user_commands.get(user_id, {})
    
    if cmd_name not in cmds:
        await callback.answer("Команда не найдена", show_alert=True)
        return
    
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Удалить",
        callback_data="del_cmd_" + cmd_name,
        style="danger",
        icon_custom_emoji_id=eid("trash")
    )
    builder.button(
        text="Назад",
        callback_data="my_commands",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(1)
    
    text = (
        em('code') + " <b>Команда:</b> <code>" + cmd_name + "</code>\n"
        + em('file') + " <b>Ответ:</b>\n<blockquote>" + cmds[cmd_name][:200] + "</blockquote>"
    )
    await safe_edit(callback.message, text, reply_markup=builder.as_markup())
    await callback.answer()

@dp.callback_query(F.data.startswith("del_cmd_"))
async def delete_command(callback: types.CallbackQuery):
    """Удаление команды"""
    cmd_name = callback.data.replace("del_cmd_", "")
    user_id = callback.from_user.id
    
    if user_id in user_commands and cmd_name in user_commands[user_id]:
        del user_commands[user_id][cmd_name]
        await callback.answer("Команда удалена!", show_alert=True)
    else:
        await callback.answer("Команда не найдена", show_alert=True)
    
    await my_commands_list(callback)

# ============ СОХРАНЕНИЕ/ЗАГРУЗКА КОНФИГОВ ============

@dp.callback_query(F.data == "save_config")
async def save_config(callback: types.CallbackQuery, state: FSMContext):
    """Сохранение конфигурации команд"""
    user_id = callback.from_user.id
    cmds = user_commands.get(user_id, {})
    
    if not cmds:
        await callback.answer("Нет команд для сохранения!", show_alert=True)
        return
    
    text = (
        em('write') + " <b>Введите название конфига:</b>\n\n"
        + em('info') + " Конфиг сохранит все команды и настройки приветствия"
    )
    await safe_edit(callback.message, text, reply_markup=get_back_keyboard("commands_menu"))
    await state.set_state(ConfigStates.waiting_for_config_name)
    await callback.answer()

@dp.message(ConfigStates.waiting_for_config_name)
async def config_get_name(message: types.Message, state: FSMContext):
    """Сохранение конфига с указанным именем"""
    config_name = message.text.strip()
    user_id = message.from_user.id
    
    if not config_name:
        await safe_send(message.chat.id, em('cross') + " Введите название конфига")
        return
    
    if user_id not in saved_configs:
        saved_configs[user_id] = {}
    
    saved_configs[user_id][config_name] = {
        "commands": user_commands.get(user_id, {}).copy(),
        "welcome_enabled": user_welcome_enabled.get(user_id, False),
    }
    
    await state.clear()
    
    text = (
        em('check') + " <b>Конфиг '" + config_name + "' успешно сохранен!</b>\n"
        + em('info') + " Вы можете загрузить его на другом аккаунте"
    )
    await safe_send(message.chat.id, text, reply_markup=get_commands_keyboard())

@dp.callback_query(F.data == "load_config")
async def load_config(callback: types.CallbackQuery):
    """Загрузка сохраненного конфига"""
    user_id = callback.from_user.id
    configs = saved_configs.get(user_id, {})
    
    if not configs:
        await callback.answer("Нет сохраненных конфигов!", show_alert=True)
        return
    
    builder = InlineKeyboardBuilder()
    for cfg_name in configs:
        builder.button(text=cfg_name, callback_data="ldcfg_" + cfg_name, style="default")
    builder.button(
        text="Назад",
        callback_data="commands_menu",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(2)
    
    text = (
        em('box') + " <b>Загрузить конфиг:</b>\n\n"
        + em('info') + " Выберите конфиг для загрузки\n"
        + em('info') + " Текущие команды будут заменены"
    )
    await safe_edit(callback.message, text, reply_markup=builder.as_markup())
    await callback.answer()

@dp.callback_query(F.data.startswith("ldcfg_"))
async def load_config_confirm(callback: types.CallbackQuery):
    """Подтверждение загрузки конфига"""
    config_name = callback.data.replace("ldcfg_", "")
    user_id = callback.from_user.id
    configs = saved_configs.get(user_id, {})
    
    if config_name not in configs:
        await callback.answer("Конфиг не найден!", show_alert=True)
        return
    
    cfg = configs[config_name]
    
    # Загружаем команды
    user_commands[user_id] = cfg.get("commands", {}).copy()
    
    # Загружаем настройку приветствия
    user_welcome_enabled[user_id] = cfg.get("welcome_enabled", False)
    
    # Перезапускаем обработчики
    await setup_command_handler(user_id)
    if user_welcome_enabled.get(user_id):
        await setup_welcome_handler(user_id)
    
    await callback.answer("Конфиг '" + config_name + "' загружен!", show_alert=True)
    await commands_menu(callback)

# ============ ПЛАГИНЫ ============

@dp.callback_query(F.data == "plugins_menu")
async def plugins_menu(callback: types.CallbackQuery):
    """Меню плагинов"""
    user_id = callback.from_user.id
    plugins = loaded_plugins.get(user_id, {})
    active_count = sum(1 for p in plugins.values() if p.get("active"))
    
    text = (
        em('apps') + " <b>Плагины</b>\n\n"
        + em('box') + " Загружено: " + str(len(plugins)) + "\n"
        + em('check') + " Активных: " + str(active_count) + "\n\n"
        + em('info') + " Загружайте свои Python плагины\n"
        + em('info') + " Поддерживаются: .py, .txt, .pdf, .docx, .zip\n"
        + em('info') + " Документация в канале: " + DONATION_CHANNEL
    )
    await safe_edit(callback.message, text, reply_markup=get_plugins_menu_keyboard())
    await callback.answer()

@dp.callback_query(F.data == "upload_plugin")
async def upload_plugin(callback: types.CallbackQuery, state: FSMContext):
    """Загрузка плагина"""
    text = (
        em('send') + " <b>Загрузка плагина</b>\n\n"
        + em('write') + " Отправьте файл с кодом плагина:\n"
        + em('file') + " Поддерживаемые форматы:\n"
        "• .py - Python файл\n"
        "• .txt - Текстовый файл с кодом\n"
        "• .pdf - PDF документ\n"
        "• .docx - Word документ\n"
        "• .zip - Архив с файлами\n"
        "• .tar.gz - Архив с файлами\n\n"
        + em('info') + " Бот автоматически извлечет код"
    )
    await safe_edit(callback.message, text, reply_markup=get_back_keyboard("plugins_menu"))
    await state.set_state(PluginStates.waiting_for_plugin_file)
    await callback.answer()

@dp.message(PluginStates.waiting_for_plugin_file, F.document)
async def plugin_file_received(message: types.Message, state: FSMContext):
    """Получение файла плагина"""
    document = message.document
    if not document:
        await safe_send(message.chat.id, em('cross') + " Отправьте файл!")
        return
    
    file_name = document.file_name or "plugin.py"
    
    status_msg = await safe_send(
        message.chat.id,
        em('loading') + " <b>Загружаю и анализирую файл...</b>"
    )
    
    try:
        # Скачиваем файл
        file_data = await bot.download(document)
        if hasattr(file_data, 'read'):
            file_bytes = file_data.read()
        else:
            file_bytes = file_data
        
        # Извлекаем код
        code = await extract_code_from_file(file_bytes, file_name)
        
        if not code:
            try:
                await status_msg.delete()
            except:
                pass
            await safe_send(
                message.chat.id,
                em('cross') + " <b>Не удалось извлечь Python код из файла!</b>\n"
                "Убедитесь что файл содержит код на Python.",
                reply_markup=get_plugins_menu_keyboard()
            )
            await state.clear()
            return
        
        # Валидируем код
        is_valid, validation_msg = validate_plugin_code(code)
        if not is_valid:
            try:
                await status_msg.delete()
            except:
                pass
            text = (
                em('cross') + " <b>Плагин не прошел валидацию!</b>\n\n"
                + em('info') + " " + validation_msg + "\n\n"
                + em('info') + " Плагин должен содержать:\n"
                "<code>class Plugin:</code>\n"
                "<code>async def start(self):</code>\n"
                "<code>async def stop(self):</code>"
            )
            await safe_send(message.chat.id, text, reply_markup=get_plugins_menu_keyboard())
            await state.clear()
            return
        
        await state.update_data(plugin_code=code, plugin_filename=file_name)
        
        try:
            await status_msg.delete()
        except:
            pass
        
        # Генерируем имя плагина из имени файла
        suggested_name = re.sub(r'[^a-zA-Z0-9_]', '_', file_name.rsplit('.', 1)[0])[:30]
        
        text = (
            em('check') + " <b>Код успешно извлечен!</b>\n\n"
            + em('file') + " Файл: " + file_name + "\n"
            + em('write') + " <b>Введите название для плагина:</b>\n"
            "Или нажмите кнопку для авто-названия"
        )
        
        builder = InlineKeyboardBuilder()
        builder.button(
            text=suggested_name,
            callback_data="plg_name_" + suggested_name,
            style="primary"
        )
        builder.button(
            text="Отмена",
            callback_data="plugins_menu",
            style="default",
            icon_custom_emoji_id=eid("cross")
        )
        builder.adjust(1)
        
        await safe_send(message.chat.id, text, reply_markup=builder.as_markup())
        await state.set_state(PluginStates.waiting_for_plugin_name)
        
    except Exception as ex:
        try:
            await status_msg.delete()
        except:
            pass
        await safe_send(
            message.chat.id,
            em('cross') + " <b>Ошибка при обработке файла:</b>\n<code>" + str(ex) + "</code>",
            reply_markup=get_plugins_menu_keyboard()
        )
        await state.clear()

@dp.callback_query(F.data.startswith("plg_name_"), PluginStates.waiting_for_plugin_name)
async def plugin_name_from_button(callback: types.CallbackQuery, state: FSMContext):
    """Выбор имени плагина из кнопки"""
    plugin_name = callback.data.replace("plg_name_", "")
    await install_plugin(callback, state, plugin_name)

@dp.message(PluginStates.waiting_for_plugin_name)
async def plugin_name_from_text(message: types.Message, state: FSMContext):
    """Ввод имени плагина текстом"""
    plugin_name = message.text.strip()[:30]
    if not plugin_name:
        await safe_send(message.chat.id, em('cross') + " Введите название плагина")
        return
    await install_plugin(message, state, plugin_name)

async def install_plugin(event, state: FSMContext, plugin_name: str):
    """Установка и запуск плагина"""
    user_id = event.from_user.id if hasattr(event, 'from_user') else event.chat.id
    chat_id = event.message.chat.id if hasattr(event, 'message') else event.chat.id
    
    data = await state.get_data()
    code = data.get("plugin_code", "")
    
    if not code:
        await safe_send(chat_id, em('cross') + " Код плагина не найден")
        await state.clear()
        return
    
    try:
        # Загружаем модуль
        module = await load_plugin_module(code, plugin_name)
        
        # Создаем экземпляр плагина
        account = get_active_account(user_id)
        client = account[1] if account else None
        
        plugin_instance = module.Plugin(
            client=client,
            bot=bot,
            user_id=user_id
        )
        
        # Запускаем плагин
        await plugin_instance.start()
        
        # Сохраняем в хранилище
        if user_id not in loaded_plugins:
            loaded_plugins[user_id] = {}
        
        loaded_plugins[user_id][plugin_name] = {
            "module": module,
            "instance": plugin_instance,
            "active": True,
            "start_time": datetime.now(),
            "code": code
        }
        
        await safe_send(
            chat_id,
            em('check') + " <b>Плагин '" + plugin_name + "' успешно загружен и запущен!</b>",
            reply_markup=get_plugins_menu_keyboard()
        )
        
    except Exception as ex:
        error_text = traceback.format_exc()
        logger.error(f"Plugin install error: {error_text}")
        await safe_send(
            chat_id,
            em('cross') + " <b>Ошибка запуска плагина:</b>\n<code>" + str(ex)[:500] + "</code>",
            reply_markup=get_plugins_menu_keyboard()
        )
    
    await state.clear()

@dp.callback_query(F.data == "my_plugins")
async def my_plugins_list(callback: types.CallbackQuery):
    """Список загруженных плагинов"""
    user_id = callback.from_user.id
    plugins = loaded_plugins.get(user_id, {})
    
    if not plugins:
        builder = InlineKeyboardBuilder()
        builder.button(
            text="Назад",
            callback_data="plugins_menu",
            style="default",
            icon_custom_emoji_id=eid("back")
        )
        text = em('unlock') + " <b>Нет загруженных плагинов</b>"
        await safe_edit(callback.message, text, reply_markup=builder.as_markup())
        await callback.answer()
        return
    
    builder = InlineKeyboardBuilder()
    for pname, pdata in plugins.items():
        status = "✅" if pdata.get("active") else "⏸"
        builder.button(
            text=status + " " + pname,
            callback_data="plg_info_" + pname,
            style="default"
        )
    builder.button(
        text="Назад",
        callback_data="plugins_menu",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    builder.adjust(1)
    
    text = em('box') + " <b>Мои плагины (" + str(len(plugins)) + "):</b>"
    await safe_edit(callback.message, text, reply_markup=builder.as_markup())
    await callback.answer()

@dp.callback_query(F.data.startswith("plg_info_"))
async def plugin_info(callback: types.CallbackQuery):
    """Информация о плагине"""
    plugin_name = callback.data.replace("plg_info_", "")
    user_id = callback.from_user.id
    pdata = loaded_plugins.get(user_id, {}).get(plugin_name)
    
    if not pdata:
        await callback.answer("Плагин не найден", show_alert=True)
        return
    
    instance = pdata["instance"]
    stats = {}
    try:
        stats = instance.get_stats()
    except:
        pass
    
    uptime = str(datetime.now() - pdata["start_time"]) if pdata["start_time"] else "Не запущен"
    
    text = (
        em('apps') + " <b>Плагин: " + plugin_name + "</b>\n\n"
        + em('check') + " Статус: " + ("Активен" if pdata["active"] else "Остановлен") + "\n"
        + em('clock') + " Аптайм: " + uptime + "\n"
        + em('stats') + " Статистика: " + str(stats) + "\n\n"
        + em('info') + " Управление плагином:"
    )
    await safe_edit(
        callback.message,
        text,
        reply_markup=get_plugin_actions_keyboard(plugin_name, pdata["active"])
    )
    await callback.answer()

@dp.callback_query(F.data.startswith("plg_start_"))
async def plugin_start(callback: types.CallbackQuery):
    """Запуск плагина"""
    plugin_name = callback.data.replace("plg_start_", "")
    user_id = callback.from_user.id
    pdata = loaded_plugins.get(user_id, {}).get(plugin_name)
    
    if not pdata:
        await callback.answer("Плагин не найден", show_alert=True)
        return
    
    try:
        await pdata["instance"].start()
        pdata["active"] = True
        pdata["start_time"] = datetime.now()
        await callback.answer("Плагин запущен!", show_alert=True)
    except Exception as ex:
        await callback.answer("Ошибка запуска: " + str(ex)[:100], show_alert=True)
    
    await plugin_info(callback)

@dp.callback_query(F.data.startswith("plg_stop_"))
async def plugin_stop(callback: types.CallbackQuery):
    """Остановка плагина"""
    plugin_name = callback.data.replace("plg_stop_", "")
    user_id = callback.from_user.id
    pdata = loaded_plugins.get(user_id, {}).get(plugin_name)
    
    if not pdata:
        await callback.answer("Плагин не найден", show_alert=True)
        return
    
    try:
        await pdata["instance"].stop()
        pdata["active"] = False
        await callback.answer("Плагин остановлен!", show_alert=True)
    except Exception as ex:
        await callback.answer("Ошибка остановки: " + str(ex)[:100], show_alert=True)
    
    await plugin_info(callback)

@dp.callback_query(F.data.startswith("plg_stats_"))
async def plugin_stats(callback: types.CallbackQuery):
    """Статистика плагина"""
    plugin_name = callback.data.replace("plg_stats_", "")
    user_id = callback.from_user.id
    pdata = loaded_plugins.get(user_id, {}).get(plugin_name)
    
    if not pdata:
        await callback.answer("Плагин не найден", show_alert=True)
        return
    
    stats = {}
    try:
        stats = pdata["instance"].get_stats()
    except:
        stats = {"error": "Не удалось получить статистику"}
    
    text = (
        em('stats') + " <b>Статистика плагина '" + plugin_name + "':</b>\n\n"
        "<code>" + str(stats) + "</code>"
    )
    
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Назад",
        callback_data="plg_info_" + plugin_name,
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    await safe_edit(callback.message, text, reply_markup=builder.as_markup())
    await callback.answer()

@dp.callback_query(F.data.startswith("plg_del_"))
async def plugin_delete(callback: types.CallbackQuery):
    """Удаление плагина"""
    plugin_name = callback.data.replace("plg_del_", "")
    user_id = callback.from_user.id
    pdata = loaded_plugins.get(user_id, {}).get(plugin_name)
    
    if not pdata:
        await callback.answer("Плагин не найден", show_alert=True)
        return
    
    try:
        if pdata["active"]:
            await pdata["instance"].stop()
    except:
        pass
    
    del loaded_plugins[user_id][plugin_name]
    
    await callback.answer("Плагин '" + plugin_name + "' удален!", show_alert=True)
    await my_plugins_list(callback)

@dp.callback_query(F.data == "plugin_docs")
async def plugin_docs(callback: types.CallbackQuery):
    """Документация по плагинам"""
    text = (
        em('file') + " <b>Документация по плагинам</b>\n\n"
        + em('info') + " <b>Поддерживаемые форматы:</b>\n"
        "• .py - Python файл\n"
        "• .txt - Текстовый файл с кодом\n"
        "• .pdf - PDF документ\n"
        "• .docx - Word документ\n"
        "• .zip - Архив с файлами\n"
        "• .tar.gz - Архив с файлами\n\n"
        + em('link') + " <b>Документация:</b> " + DONATION_CHANNEL
    )
    
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Назад",
        callback_data="plugins_menu",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    await safe_edit(callback.message, text, reply_markup=builder.as_markup())
    await callback.answer()

# ============ АДМИН ПАНЕЛЬ ============

@dp.callback_query(F.data == "admin_stats")
async def admin_stats(callback: types.CallbackQuery):
    """Статистика бота (только для админа)"""
    if callback.from_user.id != ADMIN_ID:
        await callback.answer("Нет доступа", show_alert=True)
        return
    
    total_users = len(user_sessions)
    total_accounts = sum(len(accs) for accs in user_sessions.values())
    active_bcasts = sum(len(bcasts) for bcasts in active_broadcasts.values())
    active_reacts = sum(len(reacts) for reacts in active_reactions.values())
    total_plugins = sum(len(plugs) for plugs in loaded_plugins.values())
    
    text = (
        em('stats') + " <b>Статистика бота</b>\n\n"
        + em('people') + " Пользователей: " + str(total_users) + "\n"
        + em('profile') + " Аккаунтов: " + str(total_accounts) + "\n"
        + em('clock') + " Активных рассылок: " + str(active_bcasts) + "\n"
        + em('smile') + " Активных реакций: " + str(active_reacts) + "\n"
        + em('apps') + " Загруженных плагинов: " + str(total_plugins) + "\n"
        + em('megaphone') + " Поддержка: " + SUPPORT_USERNAME
    )
    
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Назад",
        callback_data="main_menu",
        style="default",
        icon_custom_emoji_id=eid("back")
    )
    await safe_edit(callback.message, text, reply_markup=builder.as_markup())
    await callback.answer()

@dp.callback_query(F.data == "admin_broadcast")
async def admin_broadcast_start(callback: types.CallbackQuery, state: FSMContext):
    """Начало рассылки всем пользователям (только для админа)"""
    if callback.from_user.id != ADMIN_ID:
        await callback.answer("Нет доступа", show_alert=True)
        return
    
    text = (
        em('megaphone') + " <b>Рассылка всем пользователям</b>\n\n"
        + em('write') + " Отправьте сообщение с поддержкой HTML и премиум эмодзи"
    )
    await safe_edit(callback.message, text, reply_markup=get_back_keyboard("main_menu"))
    await state.set_state(AdminStates.waiting_for_broadcast)
    await callback.answer()

@dp.message(AdminStates.waiting_for_broadcast)
async def admin_broadcast_get_message(message: types.Message, state: FSMContext):
    """Получение сообщения для массовой рассылки"""
    if message.from_user.id != ADMIN_ID:
        await state.clear()
        return
    
    msg_html = message.html_text if message.html_text else message.text or ""
    await state.update_data(admin_broadcast_message=msg_html)
    
    total_users = len(user_sessions)
    
    text = (
        em('send') + " <b>Подтвердите рассылку</b>\n\n"
        + em('people') + " Пользователей: " + str(total_users) + "\n"
        + em('file') + " Сообщение:\n"
        "<blockquote>" + msg_html[:200] + ("..." if len(msg_html) > 200 else "") + "</blockquote>\n\n"
        "Отправить всем пользователям?"
    )
    
    builder = InlineKeyboardBuilder()
    builder.button(
        text="Отправить",
        callback_data="confirm_admin_broadcast",
        style="primary",
        icon_custom_emoji_id=eid("send")
    )
    builder.button(
        text="Отмена",
        callback_data="cancel_admin_broadcast",
        style="danger",
        icon_custom_emoji_id=eid("cross")
    )
    builder.adjust(2)
    
    await safe_send(message.chat.id, text, reply_markup=builder.as_markup())
    await state.set_state(AdminStates.waiting_for_confirm)

@dp.callback_query(F.data == "confirm_admin_broadcast", AdminStates.waiting_for_confirm)
async def admin_broadcast_confirm(callback: types.CallbackQuery, state: FSMContext):
    """Подтверждение и выполнение массовой рассылки"""
    if callback.from_user.id != ADMIN_ID:
        await state.clear()
        return
    
    data = await state.get_data()
    message_html = data.get("admin_broadcast_message", "")
    
    if not message_html:
        await callback.answer("Сообщение не найдено", show_alert=True)
        await state.clear()
        return
    
    status_msg = await safe_send(
        callback.message.chat.id,
        em('loading') + " <b>Рассылка началась...</b>\n"
        + em('people') + " Пользователей: " + str(len(user_sessions))
    )
    
    sent = 0
    failed = 0
    
    for user_id in user_sessions:
        try:
            await safe_send(user_id, message_html)
            sent += 1
            await asyncio.sleep(0.5)  # Задержка чтобы не флудить
        except:
            failed += 1
    
    result_text = (
        em('check') + " <b>Рассылка завершена!</b>\n\n"
        + em('check') + " Отправлено: " + str(sent) + "\n"
        + em('cross') + " Ошибок: " + str(failed)
    )
    
    try:
        await status_msg.delete()
    except:
        pass
    
    await safe_edit(callback.message, result_text, reply_markup=get_admin_keyboard())
    await state.clear()
    await callback.answer("Рассылка завершена!", show_alert=True)

@dp.callback_query(F.data == "cancel_admin_broadcast", AdminStates.waiting_for_confirm)
async def admin_broadcast_cancel(callback: types.CallbackQuery, state: FSMContext):
    """Отмена массовой рассылки"""
    if callback.from_user.id != ADMIN_ID:
        await state.clear()
        return
    
    await state.clear()
    await safe_edit(
        callback.message,
        em('cross') + " <b>Рассылка отменена</b>",
        reply_markup=get_admin_keyboard()
    )
    await callback.answer()

# ============ MAIN ============

async def main():
    """Главная функция запуска бота"""
    # Создаем папку для сессий
    os.makedirs("sessions", exist_ok=True)
    
    # Удаляем вебхук и запускаем поллинг
    try:
        await bot.delete_webhook(drop_pending_updates=True)
    except:
        pass
    
    logger.info("Бот запущен! Документация по плагинам: " + DONATION_CHANNEL)
    await dp.start_polling(bot)

if __name__ == "__main__":
    try:
        asyncio.run(main())
    except KeyboardInterrupt:
        logger.info("Бот остановлен")
    except Exception as ex:
        logger.critical(f"Critical error: {ex}", exc_info=True)
